#!/usr/bin/env python3
"""Generează dosarele de achiziții pentru avizare AFIR (proiecte LEADER depuse prin GAL).

Model: dosarul „Furnizare tractor și accesorii” – UAT Comuna Bozioru.
Intrare: data/proiecte.json (o înregistrare per dosar de achiziție).
Ieșire: output/<NN_UAT>/<dosar>/*.docx

Câmpurile lipsă se tipăresc ca spații de completat, evidențiate cu galben.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
BLANK = "____________"


# ---------------------------------------------------------------- utilitare

def lei(v):
    """Format românesc: 236632.64 -> 236.632,64"""
    if v is None or v == "":
        return None
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    s = f"{v:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def g(d, *path):
    for p in path:
        if d is None:
            return None
        d = d.get(p) if isinstance(d, dict) else None
    return d if d not in ("", [], {}) else None


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(3)
    return doc


def add_runs(par, parts, size=None):
    """parts: listă de str sau (text, opțiuni). None -> spațiu de completat galben."""
    for part in parts:
        opts = {}
        if isinstance(part, tuple):
            part, opts = part
        if part is None:
            r = par.add_run(opts.get("blank", BLANK))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = par.add_run(str(part))
        r.bold = opts.get("b", False)
        r.italic = opts.get("i", False)
        if size or opts.get("size"):
            r.font.size = Pt(opts.get("size", size))
    return par


def P(doc_or_cell, *parts, align=None, size=None, after=None, before=None):
    par = doc_or_cell.add_paragraph()
    add_runs(par, parts, size)
    if align == "c":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "j":
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    if before is not None:
        par.paragraph_format.space_before = Pt(before)
    return par


def B(text):
    return (text, {"b": True})


def shade(cell, color="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(doc, header, rows, widths=None, size=9.5, header_fill="D9E2F3"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_runs(c.paragraphs[0], [B(h)], size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            parts = val if isinstance(val, list) else [val]
            add_runs(cells[i].paragraphs[0], parts, size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def header_uat(doc, d):
    u = d["uat"]
    P(doc, B(u.get("denumire") or None), after=0)
    adr = u.get("adresa")
    P(doc, adr if adr else None, after=0, size=10)
    P(doc, "Cod de înregistrare fiscală: ", u.get("cif"), after=0, size=10)
    tel, em = u.get("telefon"), u.get("email")
    P(doc, "Tel.: ", tel, "   E-mail: ", em, after=8, size=10)


def nr_data(doc):
    P(doc, "Nr. ", None, " / ", None, after=6)


def footer_no(doc, nr_doc):
    """Subsol: «Doc. NN | Fila ___» – ca în dosarul model."""
    for sec in doc.sections:
        par = sec.footer.paragraphs[0]
        par.text = ""
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = par.add_run(f"Doc. {nr_doc} | Fila ______")
        r.font.size = Pt(9)


def semnaturi(doc, d):
    o = d["persoane"]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cols = [
        ("ÎNTOCMIT,", "Consilier achiziții publice", o.get("consilier_achizitii")),
        ("VERIFICAT,", "Compartiment financiar-contabil", o.get("contabil")),
        ("APROBAT,", d["uat"].get("functie_reprezentant", "Primar"), o.get("reprezentant")),
    ]
    for cell, (a, b, n) in zip(t.rows[0].cells, cols):
        cell.text = ""
        add_runs(cell.paragraphs[0], [B(a)])
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        P(cell, b, align="c", after=0)
        P(cell, n, align="c", after=0)
        P(cell, "..............................", align="c")


def titlu_proiect(d):
    return f"„{d['proiect']['titlu']}”" if d["proiect"].get("titlu") else None


def contract_fin(d):
    return d["proiect"].get("nr_contract_finantare")


def obiect(d):
    return d["achizitie"].get("obiect")


def tip_contract(d):
    return {"servicii": "servicii", "produse": "furnizare", "lucrari": "lucrări"}[d["achizitie"]["tip"]]


def procedura_text(d):
    a = d["achizitie"]
    if a.get("procedura") == "simplificata":
        return "Procedură simplificată – art. 7 alin. (2) și art. 113 din Legea nr. 98/2016"
    alin = {"servicii": "a)", "produse": "a)", "lucrari": "b)"}[a["tip"]]
    txt = f"Achiziție directă – art. 7 alin. (5) și alin. (7) lit. {alin} din Legea nr. 98/2016"
    if a.get("cod_seap"):
        txt += f", prin catalogul electronic SEAP (cod {a['cod_seap']}" + (f"/{a['data_seap']}" if a.get("data_seap") else "") + ")"
    return txt


# ---------------------------------------------------------------- documente

def doc_fisa_naveta(d, nr):
    doc = new_doc()
    a, pr, u = d["achizitie"], d["proiect"], d["uat"]
    P(doc, B("Formular 1"), align="r")
    P(doc, B("FIȘA NAVETĂ"), align="c", size=14, after=0)
    P(doc, "pentru documentele specifice achiziției", align="c")
    P(doc, B("Secțiunea 1 – Cererea"), before=6)
    P(doc, "Numărul de înregistrare al expeditorului / data: ________ / ____________   expediat prin: ☐ poștă ☐ curier ☐ depunere directă / electronic")
    P(doc, "Numărul de înregistrare al destinatarului / data: ________ / ____________", after=6)
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "De la:", "Către:"
    for c in t.rows[0].cells:
        shade(c)
    c0, c1 = t.rows[1].cells
    c0.text = ""
    add_runs(c0.paragraphs[0], [B("Beneficiarul de fonduri publice nerambursabile:")])
    P(c0, u.get("denumire_beneficiar") or u.get("denumire"), ", județul ", u.get("judet"), ", CIF ", u.get("cif"))
    add_runs(P(c0), [B("Reprezentantul legal de proiect:")])
    P(c0, d["persoane"].get("reprezentant"), f" – {u.get('functie_reprezentant', 'Primar')}")
    c1.text = ""
    add_runs(c1.paragraphs[0], [pr.get("crfir") or None, " / ", pr.get("ojfir") or None])
    P(c1, "Director General Adjunct CRFIR / Director OJFIR: ....................................")
    P(doc, "")
    a_tip = a["tip"]
    bifa = lambda k: "☒" if a_tip == k else "☐"
    val = None
    if a.get("valoare_contract_fara_tva"):
        val = f"{lei(a['valoare_contract_fara_tva'])} lei fără TVA ({lei(round(float(a['valoare_contract_fara_tva']) * 1.21, 2))} lei cu TVA)"
    doc_prez = ("Dosarul achiziției, conform opisului anexat, cuprinzând: programul achizițiilor pentru proiect, "
                "centralizatorul achizițiilor directe anterioare, declarația privind respectarea regulilor de evitare "
                "a conflictului de interese, certificatul constatator ONRC, documentul justificativ și oferta detaliată, "
                "analiza rezonabilității prețurilor")
    rows = [
        ["Numărul contractului de finanțare", [contract_fin(d)] + ([f" (Cererea de finanțare {pr['nr_cerere_finantare']})"] if pr.get("nr_cerere_finantare") else [])],
        ["Titlul proiectului", titlu_proiect(d)],
        ["Intervenția", [pr.get("interventie") or "DR-36 LEADER", " – Planul Strategic PAC 2023-2027 – ", pr.get("gal")]],
        ["Versiunea", "☒ Cererea inițială ☐ Versiunea nr. __"],
        ["Achiziții de", f"{bifa('servicii')} Servicii {bifa('produse')} Produse {bifa('lucrari')} Lucrări"],
        ["Obiectul achiziției", [obiect(d)] + ([f" – CPV {a['cpv']}"] if a.get("cpv") else [" – CPV ", None])],
        ["Tipul procedurii", procedura_text(d)],
        ["Documentul prezentat", doc_prez],
        ["Valoarea contractului", val],
    ]
    t2 = doc.add_table(rows=0, cols=2)
    t2.style = "Table Grid"
    for k, v in rows:
        cells = t2.add_row().cells
        cells[0].text = ""
        add_runs(cells[0].paragraphs[0], [B(k)], 10)
        shade(cells[0], "F2F2F2")
        cells[1].text = ""
        add_runs(cells[1].paragraphs[0], v if isinstance(v, list) else [v], 10)
        cells[0].width, cells[1].width = Cm(5), Cm(12)
    P(doc, "")
    P(doc, "Data: ____ / ____ / ________")
    P(doc, "Nume, prenume reprezentant legal: ", d["persoane"].get("reprezentant"))
    P(doc, "Semnătura: ______________________", after=10)
    P(doc, "Termenul limită pentru răspuns (maximum 5/10 zile lucrătoare de la data primirii) – Data: ....................", size=10)
    P(doc, "de către expertul: ...................................................... (se completează de către șeful de serviciu din cadrul CRFIR/OJFIR)", size=10)
    return doc


def lista_documente(d):
    """Lista documentelor din dosar, în ordinea opisului model. (nr, denumire, sursa)
    sursa: 'generat' = inclus în acest pachet; 'atașat' = se atașează de beneficiar."""
    a, pr = d["achizitie"], d["proiect"]
    pap = pr.get("pap")
    furn = g(a, "furnizor", "denumire")
    ctr = a.get("contract_nr_data")
    of = a.get("oferte_comparative") or []
    of_txt = "; ".join(f"{o.get('furnizor')} nr. {o.get('nr_data_oferta') or '___'}" for o in of) or "ofertele de preț"
    seap = f"cod {a['cod_seap']}" + (f"/{a['data_seap']}" if a.get("data_seap") else "") if a.get("cod_seap") else "cod DA__________"
    L = [
        ("Programul achizițiilor pentru proiect (PAP)" + (f" nr. {pap}" if pap else " nr. ____/________") + ", avizat de AFIR", "atașat"),
        ("Programul anual al achizițiilor publice " + str(d.get("an", 2026)) + " – extras cu poziția achiziției (anexa la HCL / aprobare)", "atașat"),
        ("Centralizatorul achizițiilor directe anterioare", "generat"),
        ("Referat de necesitate", "generat"),
        ("Notă estimativă (fundamentarea valorii estimate)", "generat"),
        ("Analiza rezonabilității valorii estimate", "generat"),
        ("Documente care au stat la baza analizei rezonabilității – " + of_txt, "atașat"),
        ("Notă justificativă privind alegerea procedurii", "generat"),
    ]
    if a.get("procedura") == "simplificata":
        L += [
            ("Anunț de participare simplificat publicat în SEAP și documentația de atribuire (fișa de date, caiet de sarcini, formulare, model de contract)", "atașat"),
            ("Dispoziție de numire a comisiei de evaluare, declarații de confidențialitate ale membrilor, proces-verbal de deschidere, raportul procedurii și comunicările privind rezultatul", "atașat"),
            ("Contractul de " + tip_contract(d) + (f" nr. {ctr}" if ctr else " nr. ____/________") + (f" încheiat cu {furn}" if furn else ""), "atașat"),
        ]
    else:
        L += [
            (f"Achiziție directă inițiată din catalogul electronic SEAP – {seap}", "atașat"),
            ("Document justificativ – Contract de " + tip_contract(d) + (f" nr. {ctr}" if ctr else " nr. ____/________") + ", cu anexele", "atașat"),
        ]
    L += [
        ("Oferta detaliată – oferta tehnică și financiară" + (f" {furn}" if furn else ""), "atașat"),
        ("Certificat constatator ONRC" + (f" – {furn}" if furn else " al operatorului economic"), "atașat"),
        ("Declarație privind persoanele cu funcții de decizie", "atașat"),
        ("Declarație privind respectarea regulilor de evitare a conflictului de interese și Lista de verificare a conflictului de interese (anexă)", "generat"),
        ("CV-uri ale persoanelor cu funcții de decizie", "atașat"),
        ("Copii ale actelor de identitate ale persoanelor cu funcții de decizie", "atașat"),
        ("Declarații de confidențialitate și imparțialitate (reprezentant legal, consilier achiziții publice)", "generat"),
    ]
    if a["tip"] == "lucrari":
        L.insert(7, ("Devizul general / devizele pe obiect din documentația tehnico-economică (SF/DALI/PT) aprobată", "atașat"))
    return [(f"{i:02d}", n, s) for i, (n, s) in enumerate(L, 1)]


def doc_opis(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    P(doc, B("OPIS DOCUMENTE"), align="c", size=13, after=2)
    P(doc, "Dosar achiziție – ", B(f"„{obiect(d)}”" if obiect(d) else None),
      (f" – CPV {a['cpv']}" if a.get("cpv") else ""), align="c")
    P(doc, "Proiect ", titlu_proiect(d), align="c")
    P(doc, "Contract de finanțare nr. ", contract_fin(d), align="c", after=8)
    rows = []
    for n, den, src in lista_documente(d):
        rows.append([n, den + ("" if src == "generat" else " *"), "", ""])
    table(doc, ["Nr. doc.", "Denumire document", "Nr. file", "Fila (de la – la)"], rows, widths=[1.5, 11, 1.8, 2.7])
    P(doc, "")
    P(doc, "Total dosar: ____ documente, ____ file, numerotate de la 1 la ____. Fiecare filă poartă în subsol numărul documentului din opis și numărul filei.", size=10)
    P(doc, "* Documente emise de beneficiar / operatorul economic / SEAP, care se atașează în copie conformă cu originalul.", size=9)
    P(doc, "Reprezentant legal", before=10)
    P(doc, B(d["persoane"].get("reprezentant") or None), f" – {d['uat'].get('functie_reprezentant', 'Primar')}")
    P(doc, "Semnătura și ștampila ______________")
    return doc


def doc_centralizator(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    P(doc, B("CENTRALIZATORUL ACHIZIȚIILOR DIRECTE ANTERIOARE"), align="c", size=13)
    P(doc, "Proiect ", titlu_proiect(d), align="c", after=0)
    P(doc, "Contract de finanțare nr. ", contract_fin(d), align="c", after=8)
    rows, tot, tot_tva = [], 0.0, 0.0
    for i, x in enumerate(d.get("achizitii_anterioare") or [], 1):
        v = x.get("valoare_fara_tva")
        vt = x.get("valoare_cu_tva") or (round(float(v) * 1.21, 2) if v else None)
        if v:
            tot += float(v)
        if vt:
            tot_tva += float(vt)
        rows.append([str(i), [x.get("obiect"), (f" – CPV {x['cpv']}" if x.get("cpv") else "")],
                     x.get("furnizor"), x.get("document"), lei(v), lei(vt), x.get("pozitie_pap") or "", x.get("stadiu_afir") or "________"])
    if not rows:
        rows.append(["–", "Nu au fost efectuate achiziții directe anterioare în cadrul proiectului", "–", "–", "–", "–", "–", "–"])
    rows.append(["", B("TOTAL achiziții directe anterioare"), "", "", lei(tot) if tot else "0,00", lei(tot_tva) if tot_tva else "0,00", "", ""])
    table(doc, ["Nr. crt.", "Obiectul achiziției / cod CPV", "Operator economic", "Document justificativ (nr./data)",
                "Valoare fără TVA (lei)", "Valoare cu TVA (lei)", "Poziție PAP", "Stadiu avizare AFIR"], rows, size=8.5)
    P(doc, "")
    P(doc, "Achiziția care face obiectul prezentului dosar – ", f"„{obiect(d)}”" if obiect(d) else None,
      (f" (CPV {a['cpv']}" if a.get("cpv") else " (CPV ___"),
      ", ", (lei(a.get("valoare_contract_fara_tva") or a.get("valoare_estimata_fara_tva")) or None),
      " lei fără TVA) – nu este similară, din punct de vedere al scopului, naturii și obiectului, cu achizițiile directe anterioare "
      "din cadrul proiectului; nu a fost divizată în scopul evitării aplicării procedurilor prevăzute de Legea nr. 98/2016, "
      "iar valoarea cumulată se încadrează în pragurile prevăzute la art. 7 alin. (5) din lege.", align="j")
    P(doc, "")
    semnaturi(doc, d)
    return doc


def doc_referat(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    nr_data(doc)
    P(doc, B("APROBAT,"), align="r", after=0)
    P(doc, d["uat"].get("functie_reprezentant", "Primar"), align="r", after=0)
    P(doc, d["persoane"].get("reprezentant"), align="r", after=8)
    P(doc, B("REFERAT DE NECESITATE"), align="c", size=13)
    P(doc, "privind achiziția ", B(f"„{obiect(d)}”" if obiect(d) else None), align="c")
    P(doc, "în cadrul proiectului ", titlu_proiect(d), ", contract de finanțare nr. ", contract_fin(d), align="c", after=8)
    P(doc, B("1. Necesitatea achiziției. "),
      "Prin contractul de finanțare nr. ", contract_fin(d), " încheiat cu Agenția pentru Finanțarea Investițiilor Rurale, în cadrul ",
      d["proiect"].get("interventie") or "intervenției DR-36 LEADER", " – ", d["proiect"].get("gal"),
      ", ", d["uat"].get("denumire_beneficiar") or d["uat"].get("denumire"),
      " implementează proiectul ", titlu_proiect(d),
      ". Pentru realizarea obiectivelor proiectului este necesară achiziția ", f"„{obiect(d)}”" if obiect(d) else None,
      ", prevăzută în cererea de finanțare, în bugetul indicativ (", a.get("linie_bugetara") or None,
      ") și în Programul achizițiilor pentru proiect.", align="j")
    P(doc, B("2. Obiectul și cantitățile. "), align="j")
    spec = a.get("specificatii_tehnice")
    if spec:
        for line in (spec if isinstance(spec, list) else [spec]):
            P(doc, "• ", line, align="j")
    else:
        P(doc, "• ", None, " (cantitate / caracteristici minime conform cererii de finanțare)")
    P(doc, B("3. Cod CPV: "), a.get("cpv"))
    P(doc, B("4. Valoarea estimată: "), lei(a.get("valoare_estimata_fara_tva")), " lei fără TVA, stabilită conform notei estimative, în limita bugetului aprobat.", align="j")
    P(doc, B("5. Sursa de finanțare: "), "fonduri nerambursabile FEADR – PS PAC 2023-2027 (DR-36 LEADER), contract de finanțare nr. ", contract_fin(d),
      "; TVA aferentă se suportă conform contractului de finanțare.", align="j")
    P(doc, B("6. Termen estimat de realizare: "), a.get("termen") or None, ".")
    P(doc, "")
    P(doc, "Întocmit,", after=0)
    P(doc, "Compartimentul de specialitate / Consilier achiziții publice", after=0)
    P(doc, d["persoane"].get("consilier_achizitii"))
    return doc


def doc_nota_estimativa(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    nr_data(doc)
    P(doc, B("NOTĂ ESTIMATIVĂ"), align="c", size=13)
    P(doc, "privind stabilirea valorii estimate a contractului de ", tip_contract(d), " ", B(f"„{obiect(d)}”" if obiect(d) else None), align="c")
    P(doc, "proiect ", titlu_proiect(d), align="c", after=8)
    P(doc, "În conformitate cu prevederile art. 9 – 25 din Legea nr. 98/2016 și ale art. 17 din H.G. nr. 395/2016, valoarea estimată "
      "a contractului s-a stabilit pe baza bugetului indicativ aprobat prin contractul de finanțare și a ofertelor de preț care au stat la baza "
      "fundamentării cererii de finanțare, după cum urmează:", align="j")
    of = a.get("oferte_comparative") or []
    rows = [[str(i), o.get("furnizor"), o.get("nr_data_oferta"), lei(o.get("valoare_lei_fara_tva") or (o.get("valoare_fara_tva") if (o.get("moneda") or "lei").lower() in ("lei", "ron") else None)) or (f"{lei(o.get('valoare_fara_tva'))} {o.get('moneda')}" if o.get("valoare_fara_tva") else None)] for i, o in enumerate(of, 1)]
    if not rows:
        rows = [[str(i), None, None, None] for i in (1, 2, 3)]
    table(doc, ["Nr.", "Operator economic", "Ofertă nr./data", "Valoare fără TVA"], rows, widths=[1.2, 7.5, 4, 4])
    P(doc, "")
    P(doc, "Bugetul aprobat pentru această achiziție (", a.get("linie_bugetara") or None, "): ", lei(a.get("buget_aprobat_fara_tva")) or None, " lei fără TVA.", align="j")
    P(doc, B("Valoarea estimată a contractului: "), B(lei(a.get("valoare_estimata_fara_tva")) or None), B(" lei fără TVA"),
      (f" ({lei(round(float(a['valoare_estimata_fara_tva']) * 1.21, 2))} lei cu TVA)" if a.get("valoare_estimata_fara_tva") else ""), ".", align="j")
    P(doc, "Valoarea estimată nu depășește bugetul aprobat și valoarea înscrisă în Programul achizițiilor pentru proiect.", align="j")
    P(doc, "")
    semnaturi(doc, d)
    return doc


def _ofv(o):
    """Valoarea ofertei în lei (sau None)."""
    if o.get("valoare_lei_fara_tva") is not None:
        return float(o["valoare_lei_fara_tva"])
    if (o.get("moneda") or "lei").lower() in ("lei", "ron") and o.get("valoare_fara_tva") is not None:
        return float(o["valoare_fara_tva"])
    return None


def doc_analiza(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a, pr = d["achizitie"], d["proiect"]
    nr_data(doc)
    P(doc, B("ANALIZA REZONABILITĂȚII VALORII ESTIMATE"), align="c", size=13, after=0)
    P(doc, f"a contractului de achiziție publică de {tip_contract(d)}", align="c", after=0)
    P(doc, B(f"„{obiect(d)}”" if obiect(d) else None), (f" – CPV {a['cpv']}" if a.get("cpv") else ""), align="c", after=0)
    P(doc, "aferent proiectului ", titlu_proiect(d), align="c", after=0)
    P(doc, "Contract de finanțare nr. ", contract_fin(d), align="c", after=8)
    P(doc, B("I. Baza legală"))
    for t in [
        "art. 9 alin. (1), art. 7 alin. (5) și alin. (7) din Legea nr. 98/2016 privind achizițiile publice, cu modificările și completările ulterioare;",
        "art. 17 din H.G. nr. 395/2016 pentru aprobarea Normelor metodologice de aplicare a prevederilor referitoare la atribuirea contractului de achiziție publică din Legea nr. 98/2016;",
        "Anexa IV la Contractul de finanțare – Instrucțiuni privind achizițiile publice pentru beneficiarii publici PS 2023-2027, potrivit cărora prețul unui produs/serviciu/lucrări se consideră rezonabil dacă este mai mic sau egal cu prețul unui produs/serviciu/lucrări de același tip identificat în baza de date AFIR, pe internet sau obținut prin solicitarea unei oferte, informațiile obținute atașându-se analizei;",
    ]:
        P(doc, "• ", t, align="j")
    P(doc, "• Cererea de finanțare nr. ", pr.get("nr_cerere_finantare"), " și bugetul indicativ al proiectului (Anexa III.2 la Contractul de finanțare).", align="j")
    P(doc, B("II. Scopul analizei"))
    if a.get("contract_nr_data"):
        P(doc, f"Scopul analizei este verificarea caracterului rezonabil al valorii contractului de {tip_contract(d)} nr. {a['contract_nr_data']}",
          (f", atribuit prin achiziție directă din catalogul electronic SEAP (cod {a['cod_seap']})" if a.get("cod_seap") else ""),
          ", prin raportarea prețului la prețurile ofertate pentru produse/servicii/lucrări de același tip de operatori economici.", align="j")
    else:
        P(doc, "Scopul analizei este verificarea caracterului rezonabil al valorii estimate a achiziției, prin raportarea acesteia la prețurile "
          "ofertate pentru produse/servicii/lucrări de același tip de operatori economici.", align="j")
    P(doc, B("III. Documentele care au stat la baza analizei"))
    P(doc, "• Referatul de necesitate nr. ________ din ________ și Nota estimativă nr. ________ din ________ – valoare estimată ",
      lei(a.get("valoare_estimata_fara_tva")), " lei fără TVA;", align="j")
    if a.get("contract_nr_data"):
        P(doc, f"• Contractul de {tip_contract(d)} nr. {a['contract_nr_data']} încheiat cu ", g(a, "furnizor", "denumire"),
          " – ", lei(a.get("valoare_contract_fara_tva")), " lei fără TVA;", align="j")
    of = a.get("oferte_comparative") or []
    for o in of:
        P(doc, "• ", o.get("furnizor"), " – Oferta nr. ", o.get("nr_data_oferta"), ";")
    if not of:
        P(doc, "• Ofertele de preț nr. ________ (se completează cu cel puțin o ofertă / sursă de preț pentru același tip de produs/serviciu/lucrare).")
    P(doc, B("IV. Analiza comparativă (lei, fără TVA)"))
    val_ref = a.get("valoare_contract_fara_tva") or a.get("valoare_estimata_fara_tva")
    rows = []
    for o in of:
        v = _ofv(o)
        dif = None
        if v and val_ref:
            dd = float(val_ref) - v
            sign = "−" if dd < 0 else "+"
            dif = f"{sign}{lei(abs(dd))} lei ({sign}{lei(abs(dd) / v * 100)}%)"
        val_txt = lei(v) if v else (f"{lei(o['valoare_fara_tva'])} {o.get('moneda')}" if o.get("valoare_fara_tva") else None)
        rows.append([[o.get("furnizor"), " – Oferta nr. ", o.get("nr_data_oferta")], val_txt, dif])
    if not rows:
        rows = [[None, None, None] for _ in range(3)]
    table(doc, ["Operator economic / ofertă", "Valoare ofertă (lei fără TVA)", "Diferență valoare analizată față de ofertă"], rows, widths=[8.5, 4, 4.5])
    P(doc, "Valoare analizată: ", B(lei(val_ref) or None), B(" lei fără TVA"),
      (" (valoarea contractului)" if a.get("valoare_contract_fara_tva") else " (valoarea estimată)"), ".", before=4)
    vs = [_ofv(o) for o in of if _ofv(o)]
    if vs and val_ref:
        mn = min(vs)
        ok = float(val_ref) <= mn
        P(doc, "Cel mai mic preț ofertat: ", B(lei(mn)), " lei fără TVA. Rezonabil (valoare ≤ minim): ", B("DA" if ok else "NU – se justifică diferența"), ".")
    P(doc, B("V. Analiza tehnică"))
    P(doc, "Produsele/serviciile/lucrările analizate corespund specificațiilor din cererea de finanțare și au aceleași caracteristici principale cu cele ofertate",
      (f": {a['specificatii_tehnice'] if isinstance(a.get('specificatii_tehnice'), str) else '; '.join(a['specificatii_tehnice'])}" if a.get("specificatii_tehnice") else ""), ".", align="j")
    P(doc, B("VI. Concluzia analizei"))
    P(doc, "Valoarea analizată, de ", lei(val_ref) or None, " lei fără TVA, se încadrează în bugetul aprobat pentru ", a.get("linie_bugetara") or None,
      " (", lei(a.get("buget_aprobat_fara_tva")) or None, " lei fără TVA), în valoarea prevăzută în Programul achizițiilor pentru proiect"
      + (f" nr. {pr['pap']}" if pr.get("pap") else "") + " și în pragul prevăzut la art. 7 alin. (5) din Legea nr. 98/2016"
      + (" pentru achiziția directă" if a.get("procedura") != "simplificata" else "") + ".", align="j")
    P(doc, "În consecință, prețul este rezonabil, în sensul prevederilor Anexei IV la Contractul de finanțare.", align="j")
    P(doc, B("VII. Anexe"))
    for o in of:
        P(doc, "• ", o.get("furnizor"), " – Oferta nr. ", o.get("nr_data_oferta"), ";")
    if a.get("cod_seap"):
        P(doc, f"• Extrasul din catalogul electronic SEAP – achiziția directă {a['cod_seap']}" + (f"/{a['data_seap']}" if a.get("data_seap") else "") + ".")
    P(doc, "")
    semnaturi(doc, d)
    return doc


def doc_nota_justificativa(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    nr_data(doc)
    P(doc, B("APROBAT,"), align="r", after=0)
    P(doc, d["uat"].get("functie_reprezentant", "Primar"), align="r", after=0)
    P(doc, d["persoane"].get("reprezentant"), align="r", after=8)
    P(doc, B("NOTĂ JUSTIFICATIVĂ"), align="c", size=13, after=0)
    P(doc, "privind alegerea procedurii de atribuire", align="c", after=0)
    P(doc, "pentru achiziția ", B(f"„{obiect(d)}”" if obiect(d) else None), (f" – CPV {a['cpv']}" if a.get("cpv") else ""), align="c", after=8)
    val = a.get("valoare_estimata_fara_tva")
    if a.get("procedura") == "simplificata":
        P(doc, "Având în vedere valoarea estimată a contractului, de ", lei(val) or None,
          " lei fără TVA, care depășește pragul pentru achiziția directă prevăzut la art. 7 alin. (5) din Legea nr. 98/2016 și este inferioară "
          "pragurilor prevăzute la art. 7 alin. (1), procedura de atribuire aplicabilă este ", B("procedura simplificată"),
          ", conform art. 7 alin. (2) lit. b) și art. 113 din Legea nr. 98/2016 și art. 101 – 102 din H.G. nr. 395/2016, cu publicarea anunțului de participare simplificat în SEAP.", align="j")
        P(doc, "Criteriul de atribuire: ", B("prețul cel mai scăzut"), " / cel mai bun raport calitate-preț (se va preciza în fișa de date).", align="j")
    else:
        lit = "b)" if a["tip"] == "lucrari" else "a)"
        P(doc, "Având în vedere valoarea estimată a contractului, de ", lei(val) or None,
          f" lei fără TVA, inferioară pragului prevăzut la art. 7 alin. (5) din Legea nr. 98/2016 pentru {'lucrări' if a['tip']=='lucrari' else 'produse și servicii'}, "
          "precum și valoarea cumulată a achizițiilor similare din cadrul proiectului (conform centralizatorului achizițiilor directe anterioare), "
          f"modalitatea de atribuire aplicabilă este ", B("achiziția directă"),
          f", conform art. 7 alin. (5) și alin. (7) lit. {lit} din Legea nr. 98/2016 și art. 43 – 46 din H.G. nr. 395/2016.", align="j")
        P(doc, "Achiziția se va realiza prin catalogul electronic SEAP / prin publicarea unui anunț publicitar în SEAP, cu respectarea principiilor "
          "nediscriminării, tratamentului egal, recunoașterii reciproce, transparenței, proporționalității și asumării răspunderii (art. 2 alin. (2) din Legea nr. 98/2016). "
          "Achiziția nu a fost divizată în scopul evitării aplicării procedurilor prevăzute de lege.", align="j")
    P(doc, "Contractul se finanțează din fonduri nerambursabile FEADR – PS PAC 2023-2027, contract de finanțare nr. ", contract_fin(d), ".", align="j")
    P(doc, "")
    P(doc, "Întocmit,", after=0)
    P(doc, "Consilier achiziții publice", after=0)
    P(doc, d["persoane"].get("consilier_achizitii"))
    return doc


def doc_conflict(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    a = d["achizitie"]
    rep = d["persoane"].get("reprezentant")
    P(doc, B("DECLARAȚIE"), align="c", size=13, after=0)
    P(doc, "cu privire la respectarea regulilor privind evitarea conflictului de interese", align="c", after=8)
    P(doc, "Subsemnatul ", B(rep or None), ", reprezentant legal al proiectului: ", titlu_proiect(d), ", contract de finanțare nr. ", contract_fin(d),
      ", declar sub sancțiunile aplicabile infracțiunilor de fals în declarații și uz de fals, că au fost respectate regulile privind evitarea "
      "conflictului de interese așa cum sunt precizate în Capitolul II, Secțiunea 4 – Reguli de evitare a conflictului de interese, din Legea nr. 98/2016.", align="j")
    P(doc, "De asemenea, declar sub sancțiunile aplicabile infracțiunilor de fals în declarații și uz de fals, că datele înscrise de mine în fișa de "
      "verificare a conflictului de interese anexată sunt reale, așa cum au reieșit din documentele verificate la data întocmirii fișei.", align="j")
    P(doc, "Înțeleg că în cazul în care această declarație nu este conformă cu realitatea sunt pasibil de încălcarea prevederilor legislației penale privind falsul în declarații.", align="j")
    P(doc, "Nume și prenume reprezentant legal: ", rep)
    P(doc, "Semnătura: ______________________")
    P(doc, "Data: ____________")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    P(doc, "Anexa la declarația cu privire la respectarea regulilor privind evitarea conflictului de interese", size=9, align="r")
    P(doc, B("LISTA DE VERIFICARE A CONFLICTULUI DE INTERESE ȘI INCOMPATIBILITĂȚI DE PARTICIPARE*"), align="c")
    P(doc, "Beneficiar: ", d["uat"].get("denumire_beneficiar") or d["uat"].get("denumire"), after=0)
    P(doc, "Titlul proiectului: ", titlu_proiect(d), after=0)
    P(doc, "Codul proiectului: ", d["proiect"].get("nr_cerere_finantare"), " / Contract de finanțare nr. ", contract_fin(d), after=0)
    furn = g(a, "furnizor", "denumire")
    cui = g(a, "furnizor", "cui")
    P(doc, "Achiziția: ", obiect(d), " – Contract nr. ", a.get("contract_nr_data"), " – ", furn, (f", CUI {cui}" if cui else ""), after=6)
    simpl = a.get("procedura") == "simplificata"
    obs8 = "" if simpl else "Ofertant unic – achiziție directă"
    lucr = a["tip"] == "lucrari"
    rows = [
        [B("Comisia de evaluare"), "", ""],
        ["1. Persoanele care participă la evaluarea ofertelor sunt angajați sau dețin părți sociale, părți de interes, acțiuni din capitalul subscris al unuia dintre ofertanți, terți susținători sau subcontractanți propuși.", "NU", ""],
        ["2. Persoanele care participă la evaluarea ofertelor fac parte din consiliul de administrație/ organul de conducere sau supervizare a unuia dintre ofertanți, terți susținători sau subcontractanți propuși.", "NU", ""],
        ["3. Persoanele care participă la evaluarea ofertelor sunt soț/soție, rudă sau afin până la gradul al doilea inclusiv, cu persoane care fac parte din consiliul de administrație/organul de conducere sau de supervizare a unuia dintre ofertanți, terți susținători sau subcontractanți propuși.", "NU", ""],
        ["4. Despre persoanele care participă la evaluarea ofertelor se constată sau există indicii rezonabile/informații concrete că poate avea, direct ori indirect, un interes personal, financiar, economic sau de altă natură, ori se află într-o altă situație de natură să îi afecteze independența și imparțialitatea pe parcursul procesului de evaluare.", "NU", ""],
        [B("Ofertanții participanți"), "", ""],
        ["5. Reprezentantul legal de proiect este acționar/asociat/administrator/angajat la firmele participante la procedura de atribuire sau are legături/relații profesionale/de rudenie cu persoane cu funcții de decizie ale acestora.", "NU", ""],
        ["6. Ofertantul / ofertantul asociat/ subcontractantul/ terțul susținător are drept membri în cadrul consiliului de administrație/organului de conducere sau de supervizare şi/sau are acționari ori asociați semnificativi persoane care sunt soț/soție, rudă sau afin până la gradul al doilea inclusiv ori care se află în relații comerciale cu persoane cu funcții de decizie ale beneficiarului sau al furnizorului de servicii de achiziție implicat în procedura de atribuire.", "NU", ""],
        ["7. Ofertantul a nominalizat printre principalele persoane desemnate pentru executarea contractului persoane care sunt soț/soție, rudă sau afin până la gradul al doilea inclusiv ori care se află în relații comerciale cu persoane cu funcții de decizie ale beneficiarului sau al furnizorului de servicii de achiziție implicat în procedura de atribuire.", "NU", ""],
        ["8. Există relații/legături între firmele participante (inclusiv asociați/ subcontractanți) sau alte părți implicate în procedura de ofertare la momentul depunerii ofertei.", "NU", obs8],
        [B("Situații potențial generatoare de conflict de interese, conform Legii nr. 10/1995, republicată"), "", ""],
        ["9. Verificatorul de proiect a participat la elaborarea proiectului sau la elaborarea raportului de expertiză tehnică?", "NU" if lucr else "N/A", "" if lucr else "Nu este cazul"],
        ["10. Verificatorul de proiect este acționar majoritar/asociat/administrator sau are legături de rudenie/profesionale cu persoane cu funcții de decizie din cadrul proiectarea.", "NU" if lucr else "N/A", "" if lucr else "Nu este cazul"],
        ["11. Dirigintele de șantier este salariat/acționar majoritar/ asociat / administrator la firma care execută lucrarea.", "NU" if lucr else "N/A", "" if lucr else "Nu este cazul"],
    ]
    table(doc, ["Situații potențial generatoare de conflict de interese, conform art. 60 din Legea nr. 98/2016", "Există conflict de interese? DA/NU", "Observații"],
          rows, widths=[11.5, 2.5, 3], size=8.5)
    P(doc, "*Lista de verificare a conflictului de interese se completează după finalizarea procedurii de atribuire a contractului.", size=8.5)
    P(doc, "Nume și prenume reprezentant: ", rep)
    P(doc, "Semnătură: ______________________")
    P(doc, "Data: ____________")
    return doc


def _decl_conf(doc, d, nume, calitate):
    P(doc, B("DECLARAȚIE*"), align="c", size=13, after=0)
    P(doc, "de confidențialitate și imparțialitate", align="c", after=0)
    P(doc, "privind achiziția ", f"„{obiect(d)}”" if obiect(d) else None, " – proiect ", titlu_proiect(d), align="c", after=8)
    P(doc, "Subsemnatul ", B(nume or None), ", născut la data de ____________, în ____________________, domiciliat în "
      "_____________________________________________, posesor al actului de identitate ____ seria ____, nr. __________, CNP _____________, "
      f"în calitate de {calitate} în cadrul ", d["uat"].get("denumire_beneficiar") or d["uat"].get("denumire"),
      ", declar pe proprie răspundere, sub sancțiunea falsului în declarații, următoarele:", align="j")
    for t in [
        "a) nu dețin părți sociale, părți de interes, acțiuni din capitalul subscris al unuia dintre ofertanți/candidați, terți susținători sau subcontractanți;",
        "b) nu fac parte din consiliul de administrație/organul de conducere sau de supervizare a unuia dintre ofertanți/candidați, terți susținători ori subcontractanți propuși;",
        "c) nu am calitatea de soț/soție, rudă sau afin, până la gradul al doilea inclusiv, cu persoane care fac parte din consiliul de administrație/organul de conducere a unuia dintre ofertanți/candidați, terți susținători ori subcontractanți propuși;",
        "d) nu am niciun interes personal, financiar, economic sau de altă natură și nu mă aflu într-o altă situație de natură să-mi afecteze independența și imparțialitatea pe parcursul procesului de evaluare a ofertelor.",
    ]:
        P(doc, t, align="j")
    P(doc, "Confirm că, în situația în care aș descoperi, în cursul acțiunii de evaluare, că un astfel de interes există, voi declara imediat acest lucru și mă voi retrage din procesul de evaluare.", align="j")
    P(doc, "Totodată, mă angajez că voi păstra confidențialitatea asupra conținutului ofertelor, precum și asupra altor informații prezentate de către ofertanți/candidați, terți susținători ori subcontractanți, a căror dezvăluire ar putea aduce atingere dreptului acestora de a-și proteja proprietatea intelectuală sau secretele comerciale, precum și asupra lucrărilor de evaluare.", align="j")
    P(doc, "Înțeleg că, în cazul în care voi divulga aceste informații, sunt pasibil de încălcarea prevederilor legislației civile și penale.", align="j")
    P(doc, "Data: ____________", before=6)
    P(doc, B(nume or None), align="r", after=0)
    P(doc, "Semnătura ______________", align="r")
    P(doc, "* Se completează de către fiecare persoană implicată în procesul de achiziție.", size=9)


def doc_confidentialitate(d, nr):
    doc = new_doc()
    header_uat(doc, d)
    _decl_conf(doc, d, d["persoane"].get("reprezentant"), f"{d['uat'].get('functie_reprezentant', 'Primar')} / reprezentant legal")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    header_uat(doc, d)
    _decl_conf(doc, d, d["persoane"].get("consilier_achizitii"), "Consilier achiziții publice")
    return doc


GENERATORS = {
    "Centralizatorul achizițiilor directe anterioare": doc_centralizator,
    "Referat de necesitate": doc_referat,
    "Notă estimativă (fundamentarea valorii estimate)": doc_nota_estimativa,
    "Analiza rezonabilității valorii estimate": doc_analiza,
    "Notă justificativă privind alegerea procedurii": doc_nota_justificativa,
    "Declarație privind respectarea regulilor de evitare a conflictului de interese și Lista de verificare a conflictului de interese (anexă)": doc_conflict,
    "Declarații de confidențialitate și imparțialitate (reprezentant legal, consilier achiziții publice)": doc_confidentialitate,
}
SHORT = {
    doc_centralizator: "Centralizator achizitii directe anterioare",
    doc_referat: "Referat de necesitate",
    doc_nota_estimativa: "Nota estimativa",
    doc_analiza: "Analiza rezonabilitatii valorii estimate",
    doc_nota_justificativa: "Nota justificativa alegere procedura",
    doc_conflict: "Declaratie conflict de interese si lista de verificare",
    doc_confidentialitate: "Declaratii confidentialitate si impartialitate",
}


def slug(s):
    s = s.translate(str.maketrans("ăâîșşțţĂÂÎȘŞȚŢ", "aaissttAAISSTT"))
    return re.sub(r"[^A-Za-z0-9.\- ]+", "", s).strip()


def genereaza(d, out_root):
    folder = out_root / slug(d["folder"]) / slug(d["dosar"])
    folder.mkdir(parents=True, exist_ok=True)
    files = []

    def save(doc, nr, name):
        footer_no(doc, nr)
        p = folder / f"{nr}. {name}.docx"
        doc.save(p)
        files.append(p)

    save(doc_fisa_naveta(d, "00"), "00", "Fisa naveta - Formular 1")
    obs = d.get("observatii") or []
    if isinstance(obs, str):
        obs = [obs]
    doc = new_doc()
    P(doc, B("DE VERIFICAT ÎNAINTE DE DEPUNERE"), " (document intern – nu se include în dosar)", size=12)
    P(doc, d["uat"].get("denumire"), " – ", obiect(d), after=8)
    for o in obs:
        P(doc, "☐ ", o, align="j")
    P(doc, "☐ Completați câmpurile evidențiate cu galben (numere de înregistrare, date, valori lipsă).", align="j")
    P(doc, "☐ Atașați documentele marcate cu * în opis, numerotați filele și completați opisul.", align="j")
    p = folder / "_DE VERIFICAT.docx"
    doc.save(p)
    save(doc_opis(d, "00a"), "00a", "Opis documentatie")
    for nr, den, src in lista_documente(d):
        fn = GENERATORS.get(den)
        if fn:
            save(fn(d, nr), nr, SHORT[fn])
    return files


def main():
    data = json.loads((ROOT / "data" / "proiecte.json").read_text(encoding="utf-8"))
    out = ROOT / "output"
    only = set(sys.argv[1:])
    n = 0
    for d in data:
        if only and d["id"] not in only:
            continue
        files = genereaza(d, out)
        n += len(files)
        print(f"{d['id']}: {len(files)} documente -> {files[0].parent.relative_to(out)}")
    print(f"Total: {n} documente")


if __name__ == "__main__":
    main()
