"""Propunere financiară cu prețuri – F3, Analiza de preț, C6, C7, C8, C9, Grafic, Anexa la formular.

Prețurile resurselor vin din licitatii_bilzi/preturi/resurse.csv (baza de prețuri), consumurile pe articol din
dicționarul ANALIZE al fiecărei licitații. Foaia de calcul păstrează formulele (Resurse → Analiza → F3 → C6–C9,
Grafic, Anexa), deci schimbarea unui preț în foaia „Resurse” recalculează toată oferta. Aceleași calcule se fac și
în Python, ca să putem trece sumele în formularul de ofertă (.docx).
"""
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
RESURSE_CSV = ROOT / "preturi" / "resurse.csv"

GALBEN = PatternFill("solid", fgColor="FFFF00")
ALBASTRU = Font(color="1F4E79")
GRI = PatternFill("solid", fgColor="D9E2F3")
SUBTOTAL = PatternFill("solid", fgColor="EDEDED")
LEI = "#,##0.00"
CANT = "#,##0.000"
CATEG = ("Material", "Manopera", "Utilaj", "Transport")
CAM = 0.0225  # contribuția asiguratorie pentru muncă, % din manoperă


def citeste_resurse(path=RESURSE_CSV):
    res = {}
    with open(path, encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if not r.get("cod") or r["cod"].startswith("#"):
                continue
            r["pret_lei"] = float(r["pret_lei"].replace(",", "."))
            res[r["cod"]] = r
    return res


def calculeaza(L, resurse, adaos, indirecte):
    """Calcul numeric identic cu formulele din foaie. Întoarce dict cu PU pe categorii și totaluri."""
    an = L["analize"]
    linii, total_cat = [], dict.fromkeys(CATEG, 0.0)
    for cap, articole in L["f3"]:
        for cod, den, um, cant in articole:
            pu = dict.fromkeys(CATEG, 0.0)
            for rc, consum in an[(cap, cod)]:
                r = resurse[rc]
                pu[r["categorie"]] += consum * r["pret_lei"]
            pu = {k: round(v, 2) for k, v in pu.items()}
            val = {k: round(cant * pu[k], 2) for k in CATEG}
            for k in CATEG:
                total_cat[k] += val[k]
            linii.append((cap, cod, den, um, cant, pu, val))
    directe = sum(total_cat.values())
    cam = round(total_cat["Manopera"] * CAM, 2)
    directe_cam = directe + cam
    ind = round(directe_cam * indirecte, 2)
    profit = round((directe_cam + ind) * adaos, 2)
    total = round(directe_cam + ind + profit, 2)
    return {"linii": linii, "total_cat": total_cat, "directe": directe, "cam": cam, "indirecte": ind,
            "profit": profit, "total": total, "tva": round(total * 0.21, 2)}


def alege_adaos(L, resurse, indirecte, tinta=0.97, minim=0.05, maxim=0.25):
    """Cel mai mare adaos (pas 0,5%) între minim și maxim cu care oferta rămâne sub tinta × valoarea estimată."""
    a = maxim
    while a >= minim - 1e-9:
        c = calculeaza(L, resurse, a, indirecte)
        if c["total"] <= tinta * L["valoare"]:
            return round(a, 3), c
        a -= 0.005
    return minim, calculeaza(L, resurse, minim, indirecte)


def _antet(ws, L, firma, titlu, ncol):
    ws["A1"] = titlu
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Obiectiv: {L['titlu']} – Beneficiar: {L['autoritate_scurt']}"
    ws["A3"] = f"Ofertant: {firma['denumire']}, CUI {firma['cui']}, {firma['rc']}, Buzău"
    ws["A4"] = "Valori în lei, fără TVA"
    ws["A4"].font = Font(italic=True, size=9)


def _hdr(ws, row, cols):
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(bold=True)
        cell.fill = GRI
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _semnatura(ws, row, col, firma):
    ws.cell(row, col, "OFERTANT,")
    ws.cell(row + 1, col, firma["denumire"])
    ws.cell(row + 2, col, f"{firma['admin']} – administrator")
    ws.cell(row + 3, col, "Semnătura și ștampila")


def genereaza(L, firma, resurse, path, adaos, indirecte, calc):
    an = L["analize"]
    wb = openpyxl.Workbook()

    # ---------------- Instrucțiuni
    ins = wb.active
    ins.title = "Instructiuni"
    ins["A1"] = f"Propunere financiară – {L['titlu']} ({L['nr_anunt']})"
    ins["A1"].font = Font(bold=True, size=13)
    for i, t in enumerate([
        f"Termen: {L['termen']}, ora {L['ora']} – {L.get('depunere', '')}",
        "1. Foaia Resurse: prețurile resurselor (fără TVA) din baza de prețuri BILZI, cu sursa fiecăruia. Celulele "
        "galbene se pot modifica – oferta se recalculează.",
        "2. Foaia Analiza: consumurile pe unitatea fiecărui articol din F3 (norme uzuale, de verificat).",
        "3. F3: prețurile unitare pe material / manoperă / utilaj / transport rezultă din Analiza; CAM 2,25% din "
        "manoperă, cheltuieli indirecte și profit în celulele galbene.",
        "4. C6–C9, Grafic și Anexa se completează singure; rândurile „Verificare” trebuie să fie 0.",
        f"5. Valoarea estimată: {L['valoare']:,.2f} lei fără TVA. Oferta peste această valoare este respinsă.",
        f"Profit aplicat: {adaos * 100:.1f}% (cel mai mare pas de 0,5% cu care oferta rămâne la cel mult 97% din "
        f"valoarea estimată, între 5% și 25%).",
    ], 3):
        ins.cell(i, 1, t)
    ins.column_dimensions["A"].width = 140

    # ---------------- Resurse
    rs = wb.create_sheet("Resurse")
    _antet(rs, L, firma, "Prețuri resurse (fără TVA)", 6)
    _hdr(rs, 6, ["Cod", "Denumire resursă", "UM", "Categorie", "Preț unitar (lei)", "Sursa prețului"])
    folosite = []
    for (cap, cod), lst in an.items():
        for rc, _ in lst:
            if rc not in folosite:
                folosite.append(rc)
    folosite.sort(key=lambda c: (CATEG.index(resurse[c]["categorie"]), c))
    rrow = {}
    for i, rc in enumerate(folosite, 7):
        r = resurse[rc]
        rs.cell(i, 1, rc)
        rs.cell(i, 2, r["denumire"])
        rs.cell(i, 3, r["um"])
        rs.cell(i, 4, r["categorie"])
        c = rs.cell(i, 5, r["pret_lei"])
        c.number_format = LEI
        c.fill = GALBEN
        c.font = ALBASTRU
        rs.cell(i, 6, r["sursa"])
        rrow[rc] = i
    rlast = 6 + len(folosite)
    for col, w in zip("ABCDEF", [9, 60, 8, 11, 14, 80]):
        rs.column_dimensions[col].width = w

    # ---------------- F3 (rânduri fixe, ca Analiza să poată trimite la cantitate)
    f3 = wb.create_sheet("F3")
    _antet(f3, L, firma, "F3 – Lista cu cantitățile de lucrări – deviz ofertă", 14)
    _hdr(f3, 6, ["Nr.", "Simbol", "Denumire articol", "UM", "Cantitate", "PU material", "PU manoperă", "PU utilaj",
                 "PU transport", "Valoare material", "Valoare manoperă", "Valoare utilaj", "Valoare transport", "Total"])
    row = 7
    art_row, sub_rows = {}, []
    nr = 0
    for cap, articole in L["f3"]:
        f3.cell(row, 3, cap).font = Font(bold=True)
        row += 1
        first = row
        for cod, den, um, cant in articole:
            nr += 1
            art_row[(cap, cod)] = row
            f3.cell(row, 1, nr)
            f3.cell(row, 2, cod)
            f3.cell(row, 3, den).alignment = Alignment(wrap_text=True, vertical="top")
            f3.cell(row, 4, um)
            f3.cell(row, 5, cant).number_format = CANT
            for j, k in enumerate(CATEG):
                f3.cell(row, 6 + j, f'=ROUND(SUMIFS(Analiza!$J:$J,Analiza!$A:$A,$A{row},Analiza!$K:$K,"{k}"),2)')
                f3.cell(row, 6 + j).number_format = LEI
                f3.cell(row, 10 + j, f"=ROUND($E{row}*{get_column_letter(6 + j)}{row},2)").number_format = LEI
            f3.cell(row, 14, f"=SUM(J{row}:M{row})").number_format = LEI
            row += 1
        f3.cell(row, 3, f"Total {cap}").font = Font(bold=True)
        for j in range(10, 15):
            col = get_column_letter(j)
            f3.cell(row, j, f"=SUM({col}{first}:{col}{row - 1})").number_format = LEI
        for j in range(1, 15):
            f3.cell(row, j).fill = SUBTOTAL
        sub_rows.append((cap, row))
        row += 2
    t = row
    f3.cell(t, 3, "Total cheltuieli directe din articole").font = Font(bold=True)
    for j in range(10, 15):
        col = get_column_letter(j)
        f3.cell(t, j, "=" + "+".join(f"{col}{r}" for _, r in sub_rows)).number_format = LEI
    f3.cell(t + 1, 3, "Contribuție asiguratorie pentru muncă (CAM, % din manoperă)")
    f3.cell(t + 1, 5, CAM).number_format = "0.00%"
    f3.cell(t + 1, 11, f"=ROUND(K{t}*E{t + 1},2)").number_format = LEI
    f3.cell(t + 1, 14, f"=K{t + 1}").number_format = LEI
    f3.cell(t + 2, 3, "Total cheltuieli directe (inclusiv CAM)").font = Font(bold=True)
    f3.cell(t + 2, 14, f"=N{t}+N{t + 1}").number_format = LEI
    f3.cell(t + 3, 3, "Cheltuieli indirecte (% din cheltuieli directe)")
    f3.cell(t + 3, 5, indirecte).number_format = "0.00%"
    f3.cell(t + 3, 5).fill = GALBEN
    f3.cell(t + 3, 14, f"=ROUND(N{t + 2}*E{t + 3},2)").number_format = LEI
    f3.cell(t + 4, 3, "Profit (% din cheltuieli directe + indirecte)")
    f3.cell(t + 4, 5, adaos).number_format = "0.00%"
    f3.cell(t + 4, 5).fill = GALBEN
    f3.cell(t + 4, 14, f"=ROUND((N{t + 2}+N{t + 3})*E{t + 4},2)").number_format = LEI
    f3.cell(t + 5, 3, "TOTAL GENERAL (fără TVA) – se trece în formularul de ofertă").font = Font(bold=True)
    f3.cell(t + 5, 14, f"=ROUND(N{t + 2}+N{t + 3}+N{t + 4},2)").number_format = LEI
    f3.cell(t + 5, 14).font = Font(bold=True)
    f3.cell(t + 6, 3, "TVA")
    f3.cell(t + 6, 5, 0.21).number_format = "0%"
    f3.cell(t + 6, 14, f"=ROUND(N{t + 5}*E{t + 6},2)").number_format = LEI
    f3.cell(t + 7, 3, "TOTAL GENERAL cu TVA").font = Font(bold=True)
    f3.cell(t + 7, 14, f"=N{t + 5}+N{t + 6}").number_format = LEI
    f3.cell(t + 9, 3, "Valoarea estimată de autoritatea contractantă (fără TVA)")
    f3.cell(t + 9, 14, L["valoare"]).number_format = LEI
    f3.cell(t + 10, 3, "Încadrare în valoarea estimată")
    f3.cell(t + 10, 14, f'=IF(N{t + 5}<=N{t + 9},"DA – "&TEXT(N{t + 5}/N{t + 9},"0,0%")&" din valoarea estimată",'
                         f'"NU – depășește valoarea estimată")')
    f3.cell(t + 12, 3, "Notă: " + L["note_f3"]).alignment = Alignment(wrap_text=True)
    f3.cell(t + 12, 3).font = Font(italic=True, size=9)
    _semnatura(f3, t + 14, 3, firma)
    for col, w in zip("ABCDEFGHIJKLMN", [5, 11, 60, 6, 11, 11, 11, 10, 10, 13, 13, 12, 12, 14]):
        f3.column_dimensions[col].width = w
    f3.freeze_panes = "A7"
    tot_row = t + 5

    # ---------------- Analiza
    a = wb.create_sheet("Analiza")
    _antet(a, L, firma, "Analiza de preț pe articole – consumuri de resurse pe unitatea de măsură a articolului", 12)
    _hdr(a, 6, ["Nr. art.", "Articol", "Cod resursă", "Resursă", "UM", "Consum / UM articol", "Cantitate articol",
                "Consum total", "Preț unitar resursă (lei)", "Valoare / UM articol", "Categorie", "Observații"])
    ar = 7
    for (cap, cod), lst in an.items():
        frow = art_row[(cap, cod)]
        for rc, consum in lst:
            rr = rrow[rc]
            a.cell(ar, 1, f"=F3!A{frow}")
            a.cell(ar, 2, f"=F3!C{frow}")
            a.cell(ar, 3, rc)
            a.cell(ar, 4, f"=Resurse!B{rr}")
            a.cell(ar, 5, f"=Resurse!C{rr}")
            c = a.cell(ar, 6, consum)
            c.number_format = "0.00000"
            c.fill = GALBEN
            c.font = ALBASTRU
            a.cell(ar, 7, f"=F3!E{frow}").number_format = CANT
            a.cell(ar, 8, f"=F{ar}*G{ar}").number_format = CANT
            a.cell(ar, 9, f"=Resurse!E{rr}").number_format = LEI
            a.cell(ar, 10, f"=F{ar}*I{ar}").number_format = LEI
            a.cell(ar, 11, f"=Resurse!D{rr}")
            ar += 1
    for col, w in zip("ABCDEFGHIJKL", [7, 45, 9, 45, 7, 11, 11, 12, 12, 12, 11, 20]):
        a.column_dimensions[col].width = w
    a.freeze_panes = "A7"

    # ---------------- C6–C9
    def lista(nume, titlu, categ, col_cant, col_pu, f3_col):
        s = wb.create_sheet(nume)
        _antet(s, L, firma, titlu, 6)
        _hdr(s, 6, ["Nr.", "Denumirea resursei", "UM", col_cant, col_pu, "Valoare (lei, fără TVA)"])
        k = 7
        for rc in folosite:
            if resurse[rc]["categorie"] != categ:
                continue
            rr = rrow[rc]
            s.cell(k, 1, k - 6)
            s.cell(k, 2, f"=Resurse!B{rr}")
            s.cell(k, 3, f"=Resurse!C{rr}")
            s.cell(k, 4, f'=SUMIFS(Analiza!$H:$H,Analiza!$C:$C,"{rc}")').number_format = CANT
            s.cell(k, 5, f"=Resurse!E{rr}").number_format = LEI
            s.cell(k, 6, f"=ROUND(D{k}*E{k},2)").number_format = LEI
            k += 1
        s.cell(k, 2, "TOTAL").font = Font(bold=True)
        s.cell(k, 6, f"=SUM(F7:F{k - 1})").number_format = LEI
        s.cell(k + 1, 2, "Verificare: diferență față de F3 (rotunjiri, trebuie ≈ 0)")
        s.cell(k + 1, 6, f"=ROUND(F{k}-F3!{f3_col}{t},0)").number_format = LEI
        _semnatura(s, k + 3, 2, firma)
        for col, w in zip("ABCDEF", [5, 60, 8, 14, 14, 18]):
            s.column_dimensions[col].width = w

    lista("C6", "C6 – Lista consumurilor de resurse materiale", "Material", "Cantitate totală", "Preț unitar (lei)",
          "J")
    lista("C7", "C7 – Lista consumurilor cu mâna de lucru", "Manopera", "Total ore-om", "Tarif orar (lei)", "K")
    lista("C8", "C8 – Lista consumurilor de ore de funcționare a utilajelor de construcții", "Utilaj", "Total ore",
          "Tarif orar (lei)", "L")
    lista("C9", "C9 – Lista consumurilor privind transporturile", "Transport", "Cantitate", "Tarif (lei)", "M")

    # ---------------- Grafic
    gr = wb.create_sheet("Grafic")
    sapt = L["saptamani"]
    _antet(gr, L, firma, f"Graficul de execuție fizic și valoric – {L['durata_scurt']}", 4 + sapt)
    hdr = ["Nr.", "Categoria de lucrări", "Valoare totală"] + [f"Săpt. {w + 1} %" for w in range(sapt)] + \
          ["Verificare %"] + [f"Săpt. {w + 1} lei" for w in range(sapt)]
    _hdr(gr, 6, hdr)
    k = 7
    for i, (cap, srow) in enumerate(sub_rows):
        gr.cell(k, 1, i + 1)
        gr.cell(k, 2, cap)
        # valoarea capitolului cu CAM, indirecte și profit repartizate proporțional
        gr.cell(k, 3, f"=ROUND(F3!N{srow}/F3!N{t}*F3!N{tot_row},2)").number_format = LEI
        proc = L["grafic"][i]
        for w in range(sapt):
            c = gr.cell(k, 4 + w, proc[w] if w < len(proc) else 0)
            c.number_format = "0%"
            c.fill = GALBEN
            gr.cell(k, 5 + sapt + w, f"=ROUND($C{k}*{get_column_letter(4 + w)}{k},2)").number_format = LEI
        gr.cell(k, 4 + sapt, f"=SUM({get_column_letter(4)}{k}:{get_column_letter(3 + sapt)}{k})").number_format = "0%"
        k += 1
    gr.cell(k, 2, "TOTAL").font = Font(bold=True)
    gr.cell(k, 3, f"=SUM(C7:C{k - 1})").number_format = LEI
    for w in range(sapt):
        col = get_column_letter(5 + sapt + w)
        gr.cell(k, 5 + sapt + w, f"=SUM({col}7:{col}{k - 1})").number_format = LEI
    gr.cell(k + 1, 2, "Verificare: diferență față de totalul F3 (trebuie ≈ 0)")
    gr.cell(k + 1, 3, f"=ROUND(C{k}-F3!N{tot_row},0)").number_format = LEI
    _semnatura(gr, k + 3, 2, firma)
    gr.column_dimensions["B"].width = 50
    gr.column_dimensions["C"].width = 14

    # ---------------- Anexa la formularul de ofertă
    ax = wb.create_sheet("Anexa oferta")
    ax["A1"] = "ANEXA LA FORMULARUL DE OFERTĂ"
    ax["A1"].font = Font(bold=True, size=13)
    ax["A2"] = f"Titlul contractului: „{L['titlu']}”"
    _hdr(ax, 4, ["Denumire", "Valoare lei (fără TVA)"])
    for i, (cap, srow) in enumerate(sub_rows, 5):
        ax.cell(i, 1, cap)
        ax.cell(i, 2, f"=Grafic!C{7 + i - 5}").number_format = LEI
    k = 5 + len(sub_rows)
    ax.cell(k, 1, "TOTAL fără TVA").font = Font(bold=True)
    ax.cell(k, 2, f"=F3!N{tot_row}").number_format = LEI
    ax.cell(k + 1, 1, "TVA 21%")
    ax.cell(k + 1, 2, f"=F3!N{tot_row + 1}").number_format = LEI
    ax.cell(k + 2, 1, "TOTAL cu TVA").font = Font(bold=True)
    ax.cell(k + 2, 2, f"=F3!N{tot_row + 2}").number_format = LEI
    ax.column_dimensions["A"].width = 70
    ax.column_dimensions["B"].width = 22
    _semnatura(ax, k + 4, 1, firma)

    wb.move_sheet("F3", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)
    return path
