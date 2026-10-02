#!/usr/bin/env python3
"""Dosare de ofertă BILZI STEEL PROFILE S.R.L. – licitații acoperișuri (02.10.2026).

Model: dosarul „Lucrări acoperiș hala producție S.D.F. Onești” (ADV1550112), depus pe 01.10.2026.
Ieșire: output/*.docx și output/*.xlsx

Datele din anunț și din caietul de sarcini (nr. anunț, CPV, oră, durată, garanție, articolele și
cantitățile F3) sunt completate; prețurile și ce rămâne de verificat sunt evidențiate cu galben.
"""
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
BLANK = "____________"

FIRMA = {
    "denumire": "BILZI STEEL PROFILE S.R.L.",
    "cui": "39666285",
    "rc": "J10/757/2018",
    "sediu": "Municipiul Buzău, Șoseaua Nordului nr. 9C, județul Buzău",
    "admin": "Șerban Marius",
    "tel": "0761 144 238",
    "email": "marius807@yahoo.com",
}

# Experiență similară dovedibilă (vezi „Experiență similară BILZI STEEL PROFILE – praguri”)
EXPERIENTA = [
    ("Reabilitare și înlocuire acoperiș cămin cultural, sat Lilieci",
     "Comuna Sinești, jud. Ialomița", "nr. 2309 / 23.05.2024", "80.000,00",
     "23.05.2024 – 03.06.2024", "PV de recepție nr. 2455 / 03.06.2024"),
    ("Reparație acoperiș clădire publică Post Poliție Alexeni",
     "UAT Comuna Alexeni, jud. Ialomița", "nr. 2005 / 05.03.2024", "42.454,94",
     "05.03.2024 – 04.04.2024", "PV de recepție nr. 2453 / 04.04.2024"),
]


def B(text):
    return (text, {"b": True})


def Y(text):
    """Text propus, de verificat – evidențiat cu galben."""
    return (text, {"y": True})


LICITATII = {
    "dofteana": {
        "fisier": "Parc_Dofteana",
        "titlu": "Înlocuire acoperiș clădire Parc Dofteana – DSBC 2026 II",
        "autoritate": "REGIA NAȚIONALĂ A PĂDURILOR – ROMSILVA prin DIRECȚIA SILVICĂ BACĂU",
        "autoritate_scurt": "Direcția Silvică Bacău",
        "email_depunere": "anunturi.seap@bacau.rosilva.ro",
        "nr_anunt": "ADV1549806",
        "data_anunt": "01.10.2026",
        "termen": "07.10.2026",
        "ora": "10:00",
        "cpv": "45261910-6 – Reparare de acoperișuri",
        "valoare": 252389.69,
        "amplasament": "Clădirea Parc Dofteana, Ocolul Silvic Tg. Ocna, localitatea Dofteana, județul Bacău",
        "procedura": "cumpărare directă",
        "durata": "60 (șaizeci) de zile calendaristice de la ordinul de începere (2 luni, conform caietului de "
                  "sarcini și art. 5.1 din contract), dar nu mai târziu de 31.12.2026",
        "garantie": "24",
        "catalog": True,
        "decizie": "Pădureanu Leonard – Director, Bitir Ioan – Director Tehnic, Tabacaru Ion – Director Economic",
        "experienta": False,
        "acces": False,
        "vizita": False,
        "lucrari": [
            "demontarea învelitorii existente din țiglă ceramică, cu recuperarea materialului, a sistemului pluvial, a "
            "instalației de paratrăsnet, a șorțurilor, capacelor și tubulaturii ceramice a coșurilor, a lambriului de la "
            "intradosul streașinii și a elementelor din lemn degradate (astereală, șipci, contrașipci, pazii), precum și "
            "a firidei electrice în vederea repoziționării;",
            "reparații locale la căpriori și grinzi prin eclisare și completare de secțiune, reparații și consolidări "
            "locale la lucarne, terasă și elementele decorative din lemn;",
            "curățarea vopselei vechi de pe elementele din lemn prin sablare laser, chituire, tratament de profunzime, "
            "impregnare și vopsire cu lazură opacă microporoasă în 2 straturi;",
            "înlocuirea generală a asterelei din scândură de rășinoase de 25 mm, montarea șipcilor 50×50 mm și a "
            "contrașipcilor 30×50 mm, tratate fungicid/insecticid;",
            "montarea foliei anticondens de minimum 120 g/mp și a învelitorii noi din țiglă ceramică Creaton "
            "Melodie/Balance, roșu cupru angobat (sau echivalent), cu accesoriile de fixare, conform NP 069-2014;",
            "coamă ventilată din țiglă ceramică, dolii, șorțuri de streașină, de fronton și de racord la perete (lucarne) "
            "din tablă prevopsită de 0,50 mm și opritori de zăpadă;",
            "lambriu și pazii din lemn de larice termotratat clasa A la intradosul streașinii, prinse cu șuruburi inox;",
            "refacerea instalației de paratrăsnet (I 7-2011, SR EN 62305), cu buletin PRAM, și repoziționarea firidei "
            "electrice, de către electrician autorizat ANRE;",
            "sistem pluvial din tablă prevopsită: jgheab 125 mm, burlan Ø90 mm, colțuri, racorduri, prelungitoare, coturi, "
            "coliere și descărcare liberă la sol;",
            "coșuri de fum (4 buc., cca. 31,90 mp): reparații cu cărămidă plină și mortar CT29, refacerea rosturilor, "
            "hidrofobizare, tubulatură inox AISI 304 Ø180 mm, coronamente și șorțuri de coș din tablă prevopsită;",
            "transportul și manipularea materialelor pe șantier și evacuarea deșeurilor din demontări, colectate selectiv "
            "(OUG 92/2021, HG 856/2002).",
        ],
        "etape": [
            "Etapa 1 – Predarea amplasamentului, constatarea stării elementelor existente, mostre și fișe tehnice pentru "
            "țiglă, accesorii și tinichigerie, supuse aprobării beneficiarului înainte de montaj.",
            "Etapa 2 – Demontări pe tronsoane: învelitoare cu recuperarea țiglei, sistem pluvial, paratrăsnet, șorțuri și "
            "tubulatură coșuri, lambriu, elemente din lemn degradate, firida electrică.",
            "Etapa 3 – Verificarea structurii din lemn după decopertare; reparații prin eclisare, consolidări la lucarne și "
            "terasă (orice situație neprevăzută se comunică beneficiarului înainte de intervenție).",
            "Etapa 4 – Sablare laser, chituire, tratament și vopsire a elementelor din lemn păstrate.",
            "Etapa 5 – Astereală nouă, folie anticondens, șipci și contrașipci, țiglă ceramică, coamă ventilată.",
            "Etapa 6 – Tinichigerie (dolii, șorțuri de streașină, fronton și racord), opritori de zăpadă.",
            "Etapa 7 – Coșuri de fum: reparații zidărie, rosturi, hidrofobizare, tubulatură inox, coronamente și șorțuri.",
            "Etapa 8 – Lambriu și pazii din larice termotratat; sistem pluvial.",
            "Etapa 9 – Paratrăsnet și firida electrică; măsurători PRAM; curățenie, evacuarea deșeurilor, recepție.",
        ],
        "materiale": [
            "țiglă ceramică ", Y("Creaton Melodie/Balance, roșu cupru angobat (sau echivalent)"),
            ", cu piese speciale, coamă, capete de coamă și bandă de coamă ventilată din aceeași gamă; folie anticondens "
            "minimum 120 g/mp; tablă prevopsită de minimum 0,50 mm, finisaj mat, pentru dolii, șorțuri și coronamente; "
            "sistem pluvial din tablă prevopsită (jgheab 125 mm, burlan Ø90 mm); cherestea de rășinoase uscată tehnic și "
            "tratată; lambriu și pazii din larice termotratat clasa A; tubulatură inox AISI 304 Ø180 mm; conductor OL-Zn Ø8 mm "
            "pentru paratrăsnet.",
        ],
        "f3": [
            ("02 Demontări", [
                ("DEM001", "Demontare învelitoare țiglă ceramică existentă, cu recuperarea materialului", "mp", 600),
                ("DEM002", "Demontare sistem pluvial existent (jgheaburi, burlane, cârlige, brățări)", "ml", 175.82),
                ("DEM003", "Demontare instalație paratrăsnet existentă inclusiv console și elemente de fixare", "LS", 1),
                ("DEM004", "Demontare șorțuri, capace și elemente metalice aferente coșurilor de fum", "buc", 4),
                ("DEM005", "Demontare tubulatură ceramică existentă coș fum", "ml", 40),
                ("DEM006", "Demontare lambriu intrados streașină existent", "mp", 155),
                ("DEM007", "Demontare elemente lemn degradate (astereală, șipci, contrașipci, pazii)", "mp", 600),
                ("DEM008", "Demontare firidă electrică și elemente suport existente pentru repoziționare", "buc", 1),
                ("DEM09", "Transportul manual al materialelor, în spații libere și neaccidentate, prin purtat direct pe "
                          "primii 10 m distanță orizontală, cu încărcătura de cel mult 50 kg, la o distanță de cel mult 60 m",
                 "t", 19),
            ]),
            ("03 Reparații suport lemn", [
                ("REP001", "Reparații locale căpriori și grinzi lemn existente prin eclisare și completare secțiune în pod",
                 "mc", 1),
                ("REP002", "Reparații locale lucarne, terasă existentă, elemente decor lemn (include și consolidări "
                           "structură de lemn)", "buc", 4),
                ("REP003", "Curățare strat vechi de vopsea de pe elemente lemn prin sablare laser", "ora", 200),
                ("REP004", "Chituire, tratament de profunzime, impregnare și vopsire lemn exterior vechi după curățare laser",
                 "mp", 50),
            ]),
            ("04 Înlocuire învelitoare", [
                ("ACS001", "Înlocuire generală astereală existentă din scândură rășinoase de 25 mm", "mp", 600),
                ("ACS002", "Montaj șipci 50×50 mm și contrașipci rășinoase 30×50 mm", "mp", 600),
                ("ACS004", "Montaj folie anticondens 120 g/mp", "mp", 600),
                ("ACS005", "Montaj țiglă ceramică inclusiv accesorii de fixare", "mp", 600),
                ("ACS006", "Montaj coamă ventilată din țiglă ceramică", "ml", 92.18),
                ("ACS007", "Montaj dolie din tablă prevopsită", "ml", 26.63),
                ("ACS008", "Montaj șorț streașină din tablă prevopsită", "ml", 110),
                ("ACS009", "Montaj șorț fronton din tablă prevopsită", "ml", 9.80),
                ("ACS010", "Montaj opritori de zăpadă pentru țigla ceramică", "ml", 70),
                ("ACS012", "Montaj șorț de racord la perete (lucarne)", "ml", 25.40),
            ]),
            ("05 Montaj intrados streașină, pazii din lemn", [
                ("ACS001", "Montaj lambriu lemn – larice termotratat, clasa A", "mp", 155),
                ("ACS003", "Montaj pazie lemn – larice termotratat, clasa A", "ml", 105),
            ]),
            ("06 Lucrări la instalații electrice și de protecție la trăsnet", [
                ("ACS012", "Refacere instalație paratrăsnet acoperiș", "LS", 1),
                ("ACS013", "Repoziționare firidă electrică și adaptare suport", "buc", 1),
            ]),
            ("07 Tinichigerie și pluvial", [
                ("TIN001", "Montaj jgheab 125 mm RAL 8019", "ml", 116.01),
                ("TIN002", "Montaj burlan 90 mm / 3 m", "ml", 60),
                ("TIN003", "Montaj colț exterior jgheab 90 grade 125", "buc", 5),
                ("TIN004", "Montaj racord jgheab-burlan 125/90", "buc", 12),
                ("TIN005", "Montaj prelungitor intermediar burlan 90 (lungime produs 1 m / 1,2 m)", "buc", 12),
                ("TIN006", "Montaj coturi burlan 60 grade", "buc", 24),
                ("TIN007", "Montaj cot evacuare burlan", "buc", 12),
                ("TIN008", "Montaj coliere burlan 90", "buc", 24),
                ("TIN009", "Descărcare ape pluviale liber la sol", "buc", 12),
            ]),
            ("08 Coșuri de fum", [
                ("COS000", "Reparații și curățare exterioară coșuri de fum existente", "mp", 31.90),
                ("COS001", "Conservare finisaj original coș fum, refacere rosturi și hidrofobizare", "mp", 31.90),
                ("COS002", "Montaj tubulatură coș fum inox", "buc", 4),
                ("COS003", "Confecționare coronament coș din tablă 0,50 mm RAL 7016, ieșire 40 cm față de coș", "buc", 4),
                ("COS004", "Montaj șorțuri coș – set complet față, laterale, spate", "buc", 4),
            ]),
        ],
        "note_f3": "Sursa: F3 din documentația de atribuire (Deviz nr. 8076/1/22.07.2026, devizele 02–08). F3-ul "
                   "original are și resursele pe fiecare articol (material, manoperă, utilaj, transport); C6–C9 se "
                   "completează din aceleași liste. Culorile din F3 (jgheab RAL 8019, șorț streașină RAL 8017, "
                   "coronament RAL 7016) diferă de caietul de sarcini (roșu cupru) – de clarificat.",
    },
    "c4": {
        "fisier": "Cazarma_3589_Pav_C4",
        "titlu": "Lucrări de reparații acoperiș pav. C4, cazarma 3589 București",
        "autoritate": "INSTITUTUL NAȚIONAL DE CERCETARE-DEZVOLTARE MEDICO-MILITARĂ „CANTACUZINO”",
        "autoritate_scurt": "INCDMM „Cantacuzino”",
        "email_depunere": "office.cantacuzino@mapn.ro",
        "nr_anunt": "ADV1550471",
        "data_anunt": "30.09.2026",
        "termen": "05.10.2026",
        "ora": "14:00",
        "cpv": "45453000-7 – Lucrări de reparații generale și de renovare",
        "valoare": 77677.10,
        "amplasament": "Pavilionul C4 din cazarma 3589, INCDMM „Cantacuzino”, Splaiul Independenței nr. 103, sector 5, "
                       "București",
        "procedura": "cumpărare directă",
        "durata": "45 (patruzeci și cinci) de zile calendaristice de la data emiterii ordinului de începere",
        "garantie": "36",
        "catalog": True,
        "decizie": Y("de completat – anunțul nu le nominalizează; documentația este aprobată de Col. medic "
                     "Conf. Univ. Dr. Cătălin-Gabriel Smarandache (comandant) și avizată de Col. Mihai-Andrei "
                     "Părăușanu (director administrativ) – cereți lista la bap@cantacuzino.ro"),
        "experienta": True,
        "acces": True,
        "vizita": True,
        "lucrari": [
            "desfacerea învelitorii din olane / țigle solzi sau profilate pe șipci, inclusiv desfacerea șipcilor doliilor "
            "(480 mp), cu sortarea materialelor recuperabile;",
            "învelitoare nouă din țiglă solzi sau olane, cu coame așezate pe șipci de lemn, inclusiv doliile, paziile și "
            "șorțurile (480 mp), cu verificarea și înlocuirea foliei anticondens unde este cazul;",
            "reparații la învelitoarea din țigle profilate, cu țigle și coame în mortar de ciment, la acoperiș fără "
            "astereală (65 mp);",
            "jgheaburi din tablă zincată 0,5 mm, semirotunde D=12,5 cm, executate pe șantier, cu colțuri, capace și ștuț "
            "de racord (70 m);",
            "panouri parazăpezi din metal cu ancore prinse în căpriori (70 buc.);",
            "etanșarea / fixarea protecțiilor din tablă la îmbinări cu inele din bandă de aluminiu 1×20 mm (60 m);",
            "plasă metalică de protecție pe cadru de oțel, cu ușă din plasă (64 mp);",
            "manipularea materialelor cu macaraua (35 t) și schelă metalică de susținere 3–6 m la parter (64 mp);",
            "gestionarea deșeurilor rezultate (depozitare, sortare și evacuare).",
        ],
        "etape": [
            "Etapa 1 – Vizita la amplasament (obligatorie înainte de ofertare), predarea amplasamentului, aprobarea listei "
            "nominale de acces în cazarmă.",
            "Etapa 2 – Montarea schelei, a sistemelor de siguranță și a jgheaburilor de evacuare a molozului.",
            "Etapa 3 – Desfacerea tinichigeriei și a învelitorii, de la coamă spre streașină, pe fâșii; sortarea țiglelor "
            "recuperabile; coborârea molozului în siguranță (fără aruncare de pe acoperiș).",
            "Etapa 4 – Verificarea și repararea foliei anticondens, șipci noi, montarea țiglelor de la streașină spre coamă, "
            "cu fixare mecanică perimetrală; dolii din tablă zincată de minimum 0,5 mm.",
            "Etapa 5 – Reparațiile la zona fără astereală: înlocuirea țiglelor profilate rupte, coame în mortar M100.",
            "Etapa 6 – Jgheaburi semirotunde D=12,5 cm cu pantă de 2–5 mm/m, cârlige la 60–80 cm, îmbinări cositorite.",
            "Etapa 7 – Parazăpezi pe rândul 2–3 de țigle, ancorate în căpriori; inele de etanșare din bandă Al.",
            "Etapa 8 – Plasa metalică de protecție pe cadru de oțel, cu ușă; curățenie, evacuarea deșeurilor, recepție.",
        ],
        "materiale": [
            "țiglă solzi / olane ", Y("compatibile cu învelitoarea existentă (model și culoare de stabilit la vizită)"),
            ", coame, folie anticondens, șipci de rășinoase tratate; tablă zincată la cald de 0,5 mm pentru jgheaburi și "
            "dolii, aliaj Staniu-Plumb LP30 pentru lipituri; parazăpezi din oțel galvanizat la cald; bandă de aluminiu "
            "1×20 mm; mortar de ciment-var M100; plasă metalică și profile din oțel pentru cadru.",
        ],
        "f3": [
            ("Pavilion C4 – antemăsurătoare (Anexa 1)", [
                ("RPCT26B1", "Desfacerea învelitorilor din olane, țigle solzi sau profilate pe șipci, incl. desfacerea "
                             "șipcilor doliilor", "mp", 480),
                ("RPCI01C", "Învelitori din țiglă solzi sau olane, coame așezate pe șipci de lemn, inclusiv doliile, "
                            "paziile, șorțurile", "mp", 480),
                ("RPCI05XB", "Reparații la învelitori de țigle profilate cu țigle și coame din mortar de ciment la "
                             "acoperiș fără astereală", "mp", 65),
                ("RCSI18A", "Jgheaburi tablă zincată 0,5 mm, semirotunde, D=12,5 cm, exec. pe șantier, incl. "
                            "colțuri/capace/ștuț racord", "m", 70),
                ("RPDE13A", "Montarea panourilor de parazăpezi din metal, cu ancore", "buc", 70),
                ("RPIZD31B", "Etanșarea sau fixarea protecțiilor din tablă la îmbinări sau a izolației cu inele din "
                             "bandă Al 1×20 mm", "m", 60),
                ("W1C11B1", "Plasă metalică de protecție pe cadru de oțel, ușă din plasă – montare", "mp", 64),
                ("TRB22E6B", "Manipulat materiale și elemente prefabricate cu macara turn, greutatea sarcinii la "
                             "fiecare transport = 0,5–1 t", "tonă", 35),
                ("CB41B1", "Susțineri din schelă metalică, sarcina 1000 daN/mp, cu înălțimea de 3–6 m, la parter",
                 "mp", 64),
            ]),
        ],
        "note_f3": "Sursa: Antemăsurătoarea pav. C4 (Anexa 1, întocmit ing. Iulia Păunescu) din documentația anunțului "
                   "ADV1550471. Specificațiile tehnice (Anexa 2) cer: termen de execuție 45 de zile calendaristice, "
                   "garanție 36 de luni, vizitarea obligatorie a amplasamentului.",
    },
}


# ---------------------------------------------------------------- utilitare docx

def lei(v):
    s = f"{v:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(3)
    return doc


def add_runs(par, parts, size=None):
    """parts: str sau (text, opțiuni). None -> spațiu de completat galben."""
    for part in parts:
        opts = {}
        if isinstance(part, tuple):
            part, opts = part
        if part is None:
            r = par.add_run(opts.get("blank", BLANK))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = par.add_run(str(part))
            if opts.get("y"):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.bold = opts.get("b", False)
        r.italic = opts.get("i", False)
        if size or opts.get("size"):
            r.font.size = Pt(opts.get("size", size))
    return par


def P(doc, *parts, align=None, size=None, after=None, before=None):
    par = doc.add_paragraph()
    add_runs(par, parts, size)
    par.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                     "j": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align, par.alignment)
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    if before is not None:
        par.paragraph_format.space_before = Pt(before)
    return par


def bullet(doc, *parts):
    par = doc.add_paragraph(style="List Bullet")
    add_runs(par, parts)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def shade(cell, color="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(doc, header, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        add_runs(c.paragraphs[0], [B(h)], size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            parts = val if isinstance(val, list) else [val]
            add_runs(cells[i].paragraphs[0], parts, size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def antet(doc):
    f = FIRMA
    P(doc, "OPERATOR ECONOMIC", after=0, size=10)
    P(doc, B(f["denumire"]), after=0)
    P(doc, f"CUI {f['cui']}, {f['rc']}", after=0, size=10)
    P(doc, f["sediu"], after=10, size=10)


def semnatura(doc, data=True):
    if data:
        P(doc, "Data completării: ", None, before=12)
    P(doc, FIRMA["admin"], before=6, after=0)
    P(doc, "(administrator)", after=0)
    P(doc, "Semnătura autorizată și ștampila")


def footer(doc, text):
    for sec in doc.sections:
        par = sec.footer.paragraphs[0]
        par.text = ""
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(text + " – pagina ")
        r.font.size = Pt(8)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        run.append(t)
        fld.append(run)
        par._p.append(fld)


# ---------------------------------------------------------------- secțiuni dosar

def scrisoare(doc, L, opis):
    antet(doc)
    P(doc, "Către: ", B(L["autoritate"]) if L["autoritate"] else None, after=0)
    P(doc, "E-mail: ", L["email_depunere"], after=12)
    P(doc, B("SCRISOARE DE ÎNAINTARE"), align="c", size=13, after=10)
    parts = ["Ca urmare a anunțului de cumpărare directă ", L["nr_anunt"],
             f" publicat în SEAP la data de {L['data_anunt']}, privind atribuirea contractului „{L['titlu']}”, cod CPV ",
             L["cpv"], f", noi, {FIRMA['denumire']}, vă transmitem alăturat oferta noastră."]
    P(doc, *parts, align="j")
    if L["catalog"]:
        P(doc, "Prețul ofertat a fost publicat în catalogul electronic SEAP, sub denumirea și codul CPV din anunț, "
               f"până la data de {L['termen']}, ora ", L["ora"], ".", align="j")
    P(doc, f"Persoană de contact: {FIRMA['admin']}, telefon {FIRMA['tel']}, e-mail {FIRMA['email']}.", after=10)
    P(doc, B("OPISUL DOCUMENTELOR"), align="c", after=6)
    table(doc, ["Nr.", "Document", "Pagina"],
          [[str(i + 1), d, None] for i, d in enumerate(opis)], widths=[1.2, 13, 2.8])
    P(doc, "Data: ", None)
    semnatura(doc, data=False)


def propunere_tehnica(doc, L):
    antet(doc)
    P(doc, B("PROPUNERE TEHNICĂ"), align="c", size=13, after=0)
    P(doc, f"pentru contractul „{L['titlu']}”", align="c", after=0)
    P(doc, "Cod CPV ", L["cpv"], align="c", after=10)

    P(doc, B("1. Date generale"))
    table(doc, ["Element", "Detalii"], [
        ["Autoritate contractantă", L["autoritate"]],
        ["Ofertant", f"{FIRMA['denumire']}, CUI {FIRMA['cui']}, {FIRMA['rc']}"],
        ["Amplasament", L["amplasament"]],
        ["Obiect", [f"Execuția lucrărilor „{L['titlu']}”, conform caietului de sarcini și listei de cantități F3"]],
        ["Durata de execuție ofertată", [L["durata"]]],
        ["Începerea lucrărilor", "În maximum 72 de ore de la emiterea ordinului de începere"],
        ["Garanția lucrărilor", [L["garantie"], " de luni calendaristice de la recepția la terminarea lucrărilor"]],
        ["Distanța sediu – amplasament", [Y(L["distanta"])]],
    ], widths=[5, 12])

    P(doc, B("2. Obiectul lucrărilor"))
    P(doc, "Ne angajăm să executăm integral lucrările prevăzute în caietul de sarcini și în lista de cantități F3. "
           "Lucrările cuprind, în principal:", align="j")
    for t in L["lucrari"]:
        bullet(doc, t)
    
    P(doc, B("3. Organizarea de șantier"), before=6)
    for t in [
        "delimitarea și semnalizarea zonei de lucru și a zonei de depozitare a materialelor, stabilite împreună cu "
        "reprezentantul beneficiarului, astfel încât activitatea din clădire să fie afectată cât mai puțin;",
        "depozitarea materialelor pe paleți și suporți, ferite de umezeală și de lovituri, în ambalajul original până la montaj;",
        "schele, sisteme de acces la înălțime și protecții colective (balustrade, plase, linii de ancorare);",
        "acoperirea provizorie cu prelate a zonelor descoperite, la sfârșitul fiecărei zile și în caz de precipitații;",
        "colectarea molozului în containere sau zone delimitate și evacuarea lui periodică.",
    ]:
        bullet(doc, t)
    if L["acces"]:
        P(doc, "Accesul personalului și al vehiculelor în cazarmă se face numai pe baza listei nominale aprobate de "
               "unitate și cu respectarea regulilor interne privind accesul și protecția informațiilor.", align="j")

    P(doc, B("4. Tehnologia de execuție"), before=6)
    P(doc, "Lucrările se execută pe tronsoane, astfel încât suprafața descoperită simultan să poată fi închisă în aceeași zi.",
      align="j")
    for t in L["etape"]:
        bullet(doc, t)

    P(doc, B("5. Resurse umane"), before=6)
    table(doc, ["Funcție / meserie", "Număr", "Observații"], [
        ["Șef de echipă / conducător tehnic al lucrării", "1", "Răspunde de execuție și calitate"],
        ["Dulgheri", Y("3"), "Structură lemn, astereală, rigle"],
        ["Montatori acoperiș / tinichigii", Y("3"), "Învelitoare, tinichigerie, sistem pluvial"],
        ["Muncitori necalificați", Y("2"), "Desfaceri, manipulare materiale, curățenie"],
        ["Conducător auto", "1", "Transport materiale și moloz"],
    ], widths=[7, 2, 8])
    P(doc, "Tot personalul este instruit în domeniul securității și sănătății în muncă și pentru lucrul la înălțime.",
      align="j")

    P(doc, B("6. Utilaje, echipamente și unelte"), before=6)
    for t in ["autocamion pentru transportul materialelor și al molozului;",
              "schele metalice și scări, cu sisteme de ancorare;",
              "scule electrice profesionale: fierăstraie circulare, mașini de înșurubat, foarfeci pentru tablă, mașini de găurit;",
              "echipament pentru aplicarea soluției ignifuge;",
              "echipamente individuale de protecție: căști, hamuri, centuri de siguranță, mănuși, încălțăminte de protecție."]:
        bullet(doc, t)

    P(doc, B("7. Materiale"), before=6)
    P(doc, "Materialele respectă standardele și normele în vigoare în România și sunt însoțite de certificat de calitate "
           "sau declarație de performanță, certificat de garanție și, după caz, agrement tehnic. Principalele materiale: ",
      *L["materiale"],
      " Specificațiile care indică o anumită marcă se consideră însoțite de mențiunea „sau echivalent”.", align="j")

    P(doc, B("8. Asigurarea calității și recepția"), before=6)
    for t in ["verificarea materialelor la livrare, pe baza documentelor de însoțire;",
              "verificarea pe faze de execuție și procese-verbale de lucrări ascunse acolo unde este cazul;",
              "remedierea pe cheltuiala noastră a oricărei lucrări care nu respectă caietul de sarcini;",
              "situații de lucrări cu antemăsurătorile aferente, la finalizarea fiecărei etape."]:
        bullet(doc, t)
    P(doc, "Lucrările se execută cu respectarea Legii nr. 10/1995 privind calitatea în construcții, a HG nr. 273/1994 "
           "privind recepția lucrărilor, a normativelor NP 069-2014, C 37-1988, ST 014-1996, C 56-2002 și a celorlalte "
           "reglementări din caietul de sarcini.", align="j")

    P(doc, B("9. Securitate și sănătate în muncă, PSI și protecția mediului"), before=6)
    for t in ["respectarea Legii nr. 319/2006 a securității și sănătății în muncă, cu accent pe lucrul la înălțime;",
              "respectarea normelor PSI, inclusiv la folosirea sculelor electrice lângă materiale combustibile;",
              "limitarea prafului și a zgomotului, colectarea selectivă a deșeurilor și evacuarea lor la depozite autorizate."]:
        bullet(doc, t)

    P(doc, B("10. Grafic de execuție"), before=6)
    P(doc, "Graficul de execuție fizic și valoric este anexat la propunerea financiară.")
    P(doc, "Declarăm că propunerea tehnică respectă integral cerințele caietului de sarcini și că am ținut cont de "
           "obligațiile relevante din domeniile mediului, social și al relațiilor de muncă.", align="j")
    semnatura(doc)


def formular_oferta(doc, L):
    P(doc, B("FORMULAR DE OFERTĂ"), align="c", size=13, after=10)
    P(doc, "Către ", L["autoritate"])
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    P(doc, f"1. Examinând documentația de atribuire, subsemnații, reprezentanți ai ofertantului {FIRMA['denumire']}, "
           f"declarăm că, în conformitate cu prevederile și cerințele cuprinse în documentația mai sus menționată, vom "
           f"executa „{L['titlu']}” pentru suma de ", None, " lei fără T.V.A., la care se adaugă T.V.A. în valoare de ",
      None, " lei, în conformitate cu anexa la formularul de ofertă.", align="j")
    P(doc, "2. Ne angajăm ca, în cazul în care oferta noastră este stabilită câștigătoare, să începem lucrările cât mai "
           "curând posibil după primirea ordinului de începere și să terminăm lucrările în ",
      L["durata"], ".", align="j")
    P(doc, "3. Ne angajăm să menținem această ofertă valabilă pentru o durată de ", Y("90 (nouăzeci)"),
      " de zile, respectiv până la data de ", None, ", și ea va rămâne obligatorie pentru noi și poate fi acceptată "
      "oricând înainte de expirarea perioadei de valabilitate.", align="j")
    P(doc, "4. Am înțeles și consimțim ca, în cazul în care oferta noastră este stabilită ca fiind câștigătoare, să "
           "constituim garanția de bună execuție în conformitate cu prevederile din documentația de atribuire.", align="j")
    P(doc, "5. Precizăm că:")
    P(doc, "[ ] depunem ofertă alternativă, ale cărei detalii sunt prezentate într-un formular de ofertă separat, marcat "
           "în mod clar „alternativă”;")
    P(doc, "[X] nu depunem ofertă alternativă.")
    P(doc, "6. Până la încheierea și semnarea contractului de achiziție publică, această ofertă, împreună cu comunicarea "
           "transmisă de dumneavoastră, prin care oferta noastră este acceptată ca fiind câștigătoare, vor constitui un "
           "contract angajant între noi.", align="j")
    P(doc, "7. Înțelegem că nu sunteți obligați să acceptați oferta cu cel mai scăzut preț sau orice altă ofertă primită.",
      align="j")
    P(doc, "Data: ", None, before=8)
    P(doc, f"{FIRMA['admin']}, în calitate de administrator, legal autorizat să semnez oferta pentru și în numele "
           f"{FIRMA['denumire']}", align="j")
    P(doc, "OPERATOR ECONOMIC", after=0)
    P(doc, B(FIRMA["denumire"]), after=0)
    P(doc, "Semnătura autorizată și ștampila")


def anexa_oferta(doc, L):
    P(doc, B("ANEXA LA FORMULARUL DE OFERTĂ"), align="c", size=13, after=10)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    table(doc, ["Denumire", "Valoare lei (fără T.V.A.)"], [
        [f"{L['titlu']} – total general deviz ofertă", None],
        ["T.V.A. 21%", None],
        ["Total cu T.V.A.", None],
    ], widths=[12, 5])
    P(doc, f"Valoarea estimată de autoritatea contractantă: {lei(L['valoare'])} lei fără T.V.A. Oferta nu o poate depăși.",
      size=9)
    P(doc, "Notă: Propunerea financiară cuprinde, pe lângă Formularul de ofertă și Anexa la acesta, Lista cu cantitățile "
           "de lucrări (F3) cu prețuri unitare, listele de consumuri C6, C7, C8, C9 și Graficul de execuție fizic și valoric.",
      align="j")
    semnatura(doc)


def decl_contract(doc, L):
    antet(doc)
    P(doc, B("DECLARAȚIE"), align="c", size=13, after=0)
    P(doc, "privind însușirea modelului de contract", align="c", after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    P(doc, f"Subsemnatul {FIRMA['admin']}, reprezentant legal al {FIRMA['denumire']}, declar pe propria răspundere că, "
           f"în calitate de ofertant la procedura de {L['procedura']} organizată de ", L["autoritate"],
      f" pentru atribuirea contractului „{L['titlu']}”, ne însușim modelul de contract pus la dispoziție de autoritatea "
      "contractantă în cadrul documentației de atribuire.", align="j")
    P(doc, "În cazul în care oferta noastră va fi declarată câștigătoare, ne angajăm să semnăm contractul în forma "
           "prezentată în documentația de atribuire.", align="j")
    semnatura(doc)


def decl_art51(doc, L):
    antet(doc)
    P(doc, B("DECLARAȚIE PRIVIND SĂNĂTATEA ȘI SECURITATEA ÎN MUNCĂ"), align="c", size=13, after=0)
    P(doc, "(privind respectarea art. 51 din Legea nr. 98/2016 privind achizițiile publice)", align="c", after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    P(doc, f"Subsemnatul {FIRMA['admin']}, reprezentant împuternicit al {FIRMA['denumire']}, declar pe propria "
           "răspundere că mă angajez să realizez activitățile din contract în conformitate cu reglementările obligatorii "
           "în domeniile mediului, social și al relațiilor de muncă stabilite prin legislația adoptată la nivelul Uniunii "
           "Europene, legislația națională, prin acorduri colective sau prin tratatele, convențiile și acordurile "
           "internaționale în aceste domenii. De asemenea, declar pe propria răspundere că la elaborarea ofertei am ținut "
           "cont de obligațiile în domeniul mediului, social și al relațiilor de muncă și am inclus în ofertă costul "
           "pentru îndeplinirea acestor obligații.", align="j")
    P(doc, "Totodată, declar că am luat la cunoștință de prevederile art. 326 „Falsul în declarații” din Codul Penal.",
      align="j")
    semnatura(doc)


def decl_art60(doc, L):
    antet(doc)
    P(doc, B("DECLARAȚIE PRIVIND EVITAREA CONFLICTULUI DE INTERESE"), align="c", size=13, after=0)
    P(doc, "potrivit art. 59–60 din Legea nr. 98/2016 privind achizițiile publice", align="c", after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    P(doc, f"1. Subsemnatul {FIRMA['admin']}, reprezentant împuternicit al {FIRMA['denumire']}, cu sediul în "
           f"{FIRMA['sediu']}, în calitate de ofertant la procedura de {L['procedura']}, declar pe propria răspundere, "
           "sub sancțiunea excluderii din procedură și sub sancțiunile aplicabile faptei de fals în acte publice, că nu "
           "mă aflu în situațiile prevăzute la art. 60 din Legea nr. 98/2016.", align="j")
    P(doc, "Persoanele ce dețin funcții de decizie în cadrul autorității contractante, nominalizate potrivit art. 21 "
           "alin. (5) și (6) din H.G. nr. 395/2016, sunt: ", L["decizie"], ".", align="j")
    P(doc, "2. Voi informa imediat autoritatea contractantă dacă vor interveni modificări în prezenta declarație pe "
           "parcursul procedurii sau, în cazul în care vom fi desemnați câștigători, pe parcursul derulării contractului, "
           "având în vedere și prevederile art. 61 din Legea nr. 98/2016.", align="j")
    P(doc, "Informațiile furnizate sunt complete și corecte în fiecare detaliu. Înțeleg că autoritatea contractantă are "
           "dreptul de a solicita orice documente doveditoare și că, în cazul în care această declarație nu este conformă "
           "cu realitatea, sunt pasibil de încălcarea prevederilor legislației penale privind falsul în declarații.",
      align="j")
    semnatura(doc)


def decl_164_165_167(doc, L):
    antet(doc)
    P(doc, B("DECLARAȚIE PRIVIND NEÎNCADRAREA ÎN SITUAȚIILE"), align="c", size=13, after=0)
    P(doc, B("PREVĂZUTE LA ART. 164, 165 ȘI 167 DIN LEGEA NR. 98/2016"), align="c", size=13, after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    P(doc, f"Subsemnatul {FIRMA['admin']}, reprezentant legal al {FIRMA['denumire']}, în calitate de ofertant, declar pe "
           "propria răspundere, sub sancțiunea excluderii din procedură și a sancțiunilor aplicate faptei de fals în "
           "declarații, că:", align="j")
    for t in [
        "nici operatorul economic, nici administratorul acestuia nu au fost condamnați prin hotărâre definitivă pentru "
        "vreuna dintre infracțiunile prevăzute la art. 164 alin. (1) din Legea nr. 98/2016;",
        "operatorul economic și-a îndeplinit obligațiile de plată a impozitelor, taxelor și contribuțiilor la bugetele "
        "componente ale bugetului general consolidat, conform art. 165 din Legea nr. 98/2016;",
        "operatorul economic nu se află în niciuna dintre situațiile prevăzute la art. 167 alin. (1) din Legea nr. 98/2016 "
        "(nu este în insolvență, faliment sau lichidare, nu a încălcat obligațiile din domeniul mediului, social și al "
        "relațiilor de muncă, nu a avut deficiențe grave în executarea unor contracte anterioare și nu a prezentat "
        "informații false).",
    ]:
        bullet(doc, t)
    P(doc, "Înțeleg că autoritatea contractantă are dreptul de a solicita documente doveditoare (certificate de atestare "
           "fiscală, cazier judiciar și fiscal) în scopul verificării acestei declarații.", align="j")
    semnatura(doc)


def experienta(doc, L):
    antet(doc)
    P(doc, B("LISTA PRINCIPALELOR LUCRĂRI SIMILARE EXECUTATE"), align="c", size=13, after=0)
    P(doc, "în ultimii 5 ani", align="c", after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    table(doc, ["Nr.", "Obiectul contractului", "Beneficiar", "Contract", "Valoare fără TVA (lei)", "Perioada",
                "Document constatator"],
          [[str(i + 1), e[0], e[1], e[2], e[3], e[4], e[5]] for i, e in enumerate(EXPERIENTA)],
          widths=[0.8, 4, 3, 2.4, 2, 2.4, 2.6], size=9)
    P(doc, f"{FIRMA['denumire']} a executat lucrările de mai sus în calitate de contractant unic. Anexăm copii ale "
           "documentelor constatatoare (contract, proces-verbal de recepție la terminarea lucrărilor).", align="j")
    P(doc, B("Total: 122.454,94 lei fără TVA, în 2 contracte cu autorități publice."))
    P(doc, Y("Atenție: PV-ul Alexeni nr. 2453/04.04.2024 îl trece pe Șerban Marius ca reprezentant al „SC LAVITEX PROD SRL” "
             "(eroare materială). Cereți Primăriei Alexeni o rectificare sau o recomandare pe BILZI, altfel se poate contesta."),
      size=9)
    semnatura(doc)


def lista_personal(doc, L):
    antet(doc)
    P(doc, B("LISTA NOMINALĂ A PERSONALULUI ȘI A AUTOVEHICULELOR"), align="c", size=13, after=0)
    P(doc, "pentru accesul în cazarmă pe durata execuției lucrărilor", align="c", after=8)
    P(doc, f"Titlul contractului: „{L['titlu']}”", after=8)
    table(doc, ["Nr.", "Nume și prenume", "CNP / seria și nr. CI", "Funcția"],
          [[str(i), None, None, None] for i in range(1, 9)], widths=[1, 6, 5, 5])
    table(doc, ["Nr.", "Autovehicul (marcă, tip)", "Nr. de înmatriculare", "Conducător auto"],
          [[str(i), None, None, None] for i in range(1, 3)], widths=[1, 6, 5, 5])
    P(doc, Y("Se completează doar dacă unitatea o cere prin anunț sau la semnarea contractului."), size=9)
    semnatura(doc)


def genereaza_docx(cheie):
    L = LICITATII[cheie]
    doc = new_doc()
    opis = ["Scrisoare de înaintare și opis",
            "Certificat constatator emis de O.N.R.C. (cu CAEN 4391 autorizat)",
            "Propunerea tehnică",
            "Formularul de ofertă",
            "Anexa la formularul de ofertă",
            "Propunerea financiară: F3 cu prețuri unitare, C6, C7, C8, C9, grafic de execuție fizic și valoric",
            "Declarație privind însușirea modelului de contract",
            "Declarație privind respectarea art. 51 din Legea nr. 98/2016 (mediu, social, relații de muncă)",
            "Declarație privind evitarea conflictului de interese (art. 59–60 din Legea nr. 98/2016)"]
    sectiuni = [propunere_tehnica, formular_oferta, anexa_oferta, decl_contract, decl_art51, decl_art60]
    if cheie == "c4":
        opis += ["Declarație privind neîncadrarea în art. 164, 165 și 167 din Legea nr. 98/2016",
                 "Lista lucrărilor similare executate, cu documente constatatoare",
                 "Lista nominală a personalului și a autovehiculelor pentru acces în cazarmă"]
    if L.get("vizita"):
        opis.insert(3, "Confirmarea vizitării amplasamentului (obligatorie – fără ea oferta nu este luată în considerare)")
        sectiuni += [decl_164_165_167, experienta, lista_personal]
    scrisoare(doc, L, opis)
    for s in sectiuni:
        page_break(doc)
        s(doc, L)
    footer(doc, f"{FIRMA['denumire']} – Ofertă „{L['titlu']}”")
    OUT.mkdir(exist_ok=True)
    path = OUT / f"Dosar_oferta_BILZI_{L['fisier']}.docx"
    doc.save(path)
    compacteaza_docx(path)
    return path


def compacteaza_docx(path):
    """Scoate din șablonul python-docx părțile nefolosite (stylesWithEffects, miniatura), ca fișierul să fie mic."""
    scoase = ("word/stylesWithEffects.xml", "docProps/thumbnail.jpeg")
    with zipfile.ZipFile(path) as z:
        parti = {n: z.read(n) for n in z.namelist() if n not in scoase}
    for n in ("[Content_Types].xml", "_rels/.rels", "word/_rels/document.xml.rels"):
        x = parti[n].decode("utf-8")
        x = re.sub(r'<Override[^>]*(stylesWithEffects|thumbnail)[^>]*/>', "", x)
        x = re.sub(r'<Relationship[^>]*(stylesWithEffects|thumbnail)[^>]*/>', "", x)
        parti[n] = x.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, data in parti.items():
            z.writestr(n, data)


# ---------------------------------------------------------------- propunere financiară (xlsx)

GALBEN = PatternFill("solid", fgColor="FFFF00")
GRI = PatternFill("solid", fgColor="D9E2F3")
SUBTOTAL = PatternFill("solid", fgColor="EDEDED")
LEI = '#,##0.00'


def genereaza_xlsx(cheie):
    """F3 cu articolele și cantitățile din documentația de atribuire; prețurile unitare (galben) se completează,
    totalurile, TVA-ul și graficul de execuție se calculează din formule."""
    L = LICITATII[cheie]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F3"
    bold = Font(bold=True)
    ws["A1"] = FIRMA["denumire"]
    ws["A1"].font = bold
    ws["A2"] = f"Obiectiv: {L['titlu']} – Beneficiar: {L['autoritate_scurt']}"
    ws["A3"] = "FORMULAR F3 – Lista cu cantitățile de lucrări pe categorii de lucrări (cu prețuri)"
    ws["A3"].font = Font(bold=True, size=13)
    hdr = ["Nr.", "Cod articol", "Capitol de lucrări", "U.M.", "Cantitate", "Preț unitar fără TVA (lei)",
           "Total fără TVA (lei)"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(5, c, h)
        cell.font = bold
        cell.fill = GRI
    r = 6
    subtotaluri, capitole = [], []
    for cap, articole in L["f3"]:
        ws.cell(r, 1, cap).font = bold
        r += 1
        first = r
        for n, (cod, den, um, cant) in enumerate(articole, 1):
            ws.cell(r, 1, n)
            ws.cell(r, 2, cod)
            ws.cell(r, 3, den)
            ws.cell(r, 4, um)
            ws.cell(r, 5, cant).number_format = '#,##0.00'
            ws.cell(r, 6).fill = GALBEN
            ws.cell(r, 6).number_format = LEI
            ws.cell(r, 7, f"=ROUND(E{r}*F{r},2)").number_format = LEI
            r += 1
        ws.cell(r, 3, f"Total {cap}").font = bold
        ws.cell(r, 7, f"=SUM(G{first}:G{r - 1})").number_format = LEI
        for c in range(1, 8):
            ws.cell(r, c).fill = SUBTOTAL
        subtotaluri.append(f"G{r}")
        capitole.append((cap, f"F3!G{r}"))
        r += 2
    tot = r
    ws.cell(r, 3, "TOTAL GENERAL fără TVA").font = bold
    ws.cell(r, 7, "=" + "+".join(subtotaluri)).number_format = LEI
    ws.cell(r + 1, 3, "TVA 21%")
    ws.cell(r + 1, 7, f"=ROUND(G{r}*0.21,2)").number_format = LEI
    ws.cell(r + 2, 3, "TOTAL GENERAL cu TVA").font = bold
    ws.cell(r + 2, 7, f"=G{r}+G{r + 1}").number_format = LEI
    ws.cell(r + 4, 3, "Valoarea estimată de autoritatea contractantă (fără TVA) – oferta nu o poate depăși")
    ws.cell(r + 4, 7, L["valoare"]).number_format = LEI
    ws.cell(r + 5, 3, "Diferență față de valoarea estimată (trebuie să fie ≥ 0)")
    ws.cell(r + 5, 7, f"=G{r + 4}-G{tot}").number_format = LEI
    ws.cell(r + 7, 1, "Notă: " + L["note_f3"]).font = Font(italic=True, size=9)
    ws.cell(r + 8, 1, "Prețurile unitare includ materialele, manopera, utilajele, transportul, cheltuielile indirecte și "
                      "profitul. Se completează doar celulele galbene.").font = Font(italic=True, size=9)
    ws.cell(r + 10, 1, f"Data: ____________        {FIRMA['admin']} – administrator, semnătura și ștampila")
    for col, w in zip("ABCDEFG", [6, 11, 70, 7, 11, 16, 18]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=6, max_row=r, max_col=3):
        row[2].alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

    # Anexa la formularul de ofertă
    an = wb.create_sheet("Anexa oferta")
    an["A1"] = "ANEXA LA FORMULARUL DE OFERTĂ"
    an["A1"].font = Font(bold=True, size=13)
    an["A2"] = f"Titlul contractului: „{L['titlu']}”"
    an["A4"], an["B4"] = "Denumire", "Valoare lei (fără TVA)"
    an["A4"].font = an["B4"].font = bold
    for k, (cap, ref) in enumerate(capitole, 5):
        an.cell(k, 1, cap)
        an.cell(k, 2, f"={ref}").number_format = LEI
    k = 5 + len(capitole)
    an.cell(k, 1, "TOTAL fără TVA").font = bold
    an.cell(k, 2, f"=F3!G{tot}").number_format = LEI
    an.cell(k + 1, 1, "TVA 21%")
    an.cell(k + 1, 2, f"=F3!G{tot + 1}").number_format = LEI
    an.cell(k + 2, 1, "TOTAL cu TVA").font = bold
    an.cell(k + 2, 2, f"=F3!G{tot + 2}").number_format = LEI
    an.column_dimensions["A"].width = 70
    an.column_dimensions["B"].width = 22

    # Grafic de execuție fizic și valoric (procente pe săptămâni, de completat)
    gr = wb.create_sheet("Grafic")
    sapt = L["saptamani"]
    gr["A1"] = "GRAFIC DE EXECUȚIE FIZIC ȘI VALORIC"
    gr["A1"].font = Font(bold=True, size=13)
    gr["A2"] = f"Durata: {L['durata']}. În celulele galbene se trece procentul din capitol executat în fiecare săptămână."
    gr.cell(4, 1, "Capitol").font = bold
    gr.cell(4, 2, "Valoare (lei)").font = bold
    for w in range(sapt):
        gr.cell(4, 3 + w, f"S{w + 1}").font = bold
    gr.cell(4, 3 + sapt, "Total %").font = bold
    for k, (cap, ref) in enumerate(capitole, 5):
        gr.cell(k, 1, cap)
        gr.cell(k, 2, f"={ref}").number_format = LEI
        for w in range(sapt):
            gr.cell(k, 3 + w).fill = GALBEN
            gr.cell(k, 3 + w).number_format = '0%'
        c0 = openpyxl.utils.get_column_letter(3)
        c1 = openpyxl.utils.get_column_letter(2 + sapt)
        gr.cell(k, 3 + sapt, f"=SUM({c0}{k}:{c1}{k})").number_format = '0%'
    k = 5 + len(capitole)
    gr.cell(k, 1, "Valoare executată pe săptămână (lei)").font = bold
    for w in range(sapt):
        col = openpyxl.utils.get_column_letter(3 + w)
        gr.cell(k, 3 + w, "=" + "+".join(f"$B${i}*N({col}{i})" for i in range(5, k))).number_format = LEI
    gr.column_dimensions["A"].width = 55
    gr.column_dimensions["B"].width = 16

    OUT.mkdir(exist_ok=True)
    path = OUT / f"Propunere_financiara_BILZI_{L['fisier']}.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    LICITATII["dofteana"]["distanta"] = "cca. 190 km (Buzău – Focșani – Onești – Dofteana)"
    LICITATII["c4"]["distanta"] = "cca. 115 km (Buzău – București)"
    LICITATII["dofteana"]["saptamani"] = 9
    LICITATII["c4"]["saptamani"] = 7
    for k in LICITATII:
        print(genereaza_docx(k))
        print(genereaza_xlsx(k))
